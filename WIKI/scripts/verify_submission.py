"""Reproduce submission evidence against a disposable copy started with python main.py.

Run from any directory: python WIKI/scripts/verify_submission.py
Build dist first. Install playwright and run python -m playwright install chromium.
Optional --server-python selects a clean runtime interpreter (without Playwright).
Never uses or modifies the business database; generated credentials stay in .tmp/.
"""
import argparse
import hashlib
import json
import os
from pathlib import Path
import platform
import secrets
import shutil
import socket
import subprocess
import sys
import tempfile
import time
from urllib.request import urlopen

from openpyxl import Workbook, load_workbook
from docx import Document
from playwright.sync_api import sync_playwright, expect

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'WIKI'


def workbook(name, headers, rows):
    target = OUT / 'examples' / name
    target.parent.mkdir(exist_ok=True)
    book = Workbook()
    book.active.append(headers)
    for row in rows:
        book.active.append(row)
    book.save(target)
    return target


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--server-python', default=sys.executable)
    args = parser.parse_args()
    interpreter = str(Path(args.server_python).resolve())
    assert (ROOT / 'dist/index.html').exists(), 'Run npm run build first'
    (ROOT / '.tmp').mkdir(exist_ok=True)
    sandbox = Path(tempfile.mkdtemp(prefix='wiki-evidence-', dir=ROOT / '.tmp'))
    for name in ['main.py', 'desktop_launcher.py', 'algorithm.py', 'manage.py']:
        shutil.copyfile(ROOT / name, sandbox / name)
    for name in ['api', 'defense_scheduler', 'dist']:
        shutil.copytree(ROOT / name, sandbox / name, ignore=shutil.ignore_patterns('__pycache__'))
    with socket.socket() as probe:
        probe.bind(('127.0.0.1', 0))
        port = probe.getsockname()[1]
    base = f'http://127.0.0.1:{port}'
    password = secrets.token_urlsafe(24)
    env = {key: value for key, value in os.environ.items()
           if not key.startswith(('DJANGO_', 'DEFENSE_SCHEDULER_', 'INITIAL_ADMIN_', 'FRONTEND_DIST_'))}
    env.update(DEFENSE_SCHEDULER_PORT=str(port), DEFENSE_SCHEDULER_OPEN_BROWSER='false',
               INITIAL_ADMIN_USERNAME='wiki-reviewer', INITIAL_ADMIN_PASSWORD=password,
               PYTHONIOENCODING='utf-8', PYTHONUNBUFFERED='1')
    images = OUT / 'images'
    evidence = OUT / 'evidence'
    images.mkdir(exist_ok=True)
    evidence.mkdir(exist_ok=True)
    screenshots = []
    with (sandbox / 'startup.log').open('w', encoding='utf-8') as log:
        server = subprocess.Popen([interpreter, 'main.py'], cwd=sandbox, env=env,
                                  stdout=log, stderr=subprocess.STDOUT,
                                  creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
        try:
            for _ in range(120):
                assert server.poll() is None, f'Startup failed: {sandbox / "startup.log"}'
                try:
                    with urlopen(base + '/health/', timeout=1) as response:
                        if json.load(response)['status'] == 'ok':
                            break
                except OSError:
                    time.sleep(0.5)
            else:
                raise AssertionError('Startup timed out')
            assert (sandbox / 'app-data/INITIAL_ADMIN.txt').exists()
            teachers = workbook('teachers.xlsx', ['name', 'college', 'title', 'roles', 'availableTypes'],
                                [[f'演示教师{i}', '演示学院', '教授', '普通专家,主席,组长,秘书', '中期答辩'] for i in range(1, 8)])
            students = workbook('students.xlsx', ['name', 'studentNo', 'studentType', 'mentorName', 'campus', 'defenseTypes'],
                                [[f'演示学生{i}', f'DEMO00{i}', '学硕', '演示教师1', '创新港', '中期答辩'] for i in range(1, 7)])
            rooms = workbook('rooms.xlsx', ['campus', 'name', 'capacity', 'availableTimes'],
                             [['创新港', '演示教室101', 30, '']])
            with sync_playwright() as p:
                browser = p.chromium.launch(headless=True)
                page = browser.new_page(viewport={'width': 1440, 'height': 1050}, locale='zh-CN')
                errors = []
                page.on('pageerror', lambda error: errors.append(str(error)))

                def shot(filename):
                    page.wait_for_load_state('networkidle')
                    expect(page.locator('.el-message')).to_have_count(0, timeout=10000)
                    page.screenshot(path=str(images / filename), full_page=True)
                    screenshots.append(filename)

                page.goto(base + '/login')
                shot('20260919-01-login.png')
                page.get_by_placeholder('请输入账号').fill('wiki-reviewer')
                page.get_by_placeholder('请输入密码').fill(password)
                with page.expect_response('**/api/auth/login/') as login:
                    page.get_by_role('button', name='登录', exact=True).click()
                token = login.value.json()['data']['token']
                page.wait_for_url('**/dashboard')
                headers = {'Authorization': f'Token {token}'}
                for endpoint, path in [('teachers', teachers), ('students', students), ('rooms', rooms)]:
                    response = page.request.post(base + f'/api/{endpoint}/import_data/', headers=headers,
                                                 multipart={'file': {'name': path.name, 'mimeType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'buffer': path.read_bytes()}})
                    assert response.ok, response.text()
                    assert response.json()['success'], response.text()
                page.reload()
                shot('20260919-02-dashboard.png')
                for route, filename in [('teachers', '20260919-03-teachers.png'), ('students', '20260919-04-students.png')]:
                    page.goto(base + '/' + route)
                    shot(filename)
                page.goto(base + '/schedule-wizard?type=中期答辩&step=1')
                page.locator('.defense-choice').filter(has_text='中期答辩').click()
                page.get_by_role('button', name='下一步：准备资料').click()
                shot('20260919-05-materials.png')
                page.get_by_role('button', name='资料已核对，确认要求').click()
                form = page.locator('.rule-config-form')
                form.locator('.el-form-item').filter(has_text='每组学生人数').get_by_role('spinbutton').fill('6')
                shot('20260919-06-rules.png')
                page.get_by_role('button', name='保存要求，去生成草稿').click()
                page.get_by_role('button', name='生成排期草稿', exact=True).click()
                page.get_by_role('dialog').get_by_role('button', name='确认生成草稿').click()
                expect(page.locator('.agenda-group')).to_have_count(1, timeout=60000)
                page.locator('.agenda-members summary').click()
                shot('20260919-07-agenda.png')
                page.get_by_role('button', name='已核对，去发布与导出').click()
                page.get_by_role('button', name='校验并发布', exact=True).click()
                expect(page.locator('.wizard-finish-note')).to_contain_text('已发布，可以下载安排表了')
                shot('20260919-08-published.png')
                result = page.request.get(base + '/api/schedule/current/?defense_type=mid', headers=headers).json()
                assert result['data']['status'] == 'published'
                assert len(result['data']['groups'][0]['students']) == 6
                (evidence / 'schedule-response.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
                for endpoint, filename in [('export-excel', 'schedule.xlsx'), ('export_word', 'schedule.docx')]:
                    response = page.request.get(base + f'/api/schedule/{endpoint}/?defense_type=mid', headers=headers)
                    assert response.ok, response.status
                    (evidence / filename).write_bytes(response.body())
                book = load_workbook(evidence / 'schedule.xlsx')
                assert len(book.sheetnames) == 1
                assert '演示学生1' in str(list(book.active.values))
                doc = Document(evidence / 'schedule.docx')
                assert doc.tables
                page.goto(base + '/operation-log')
                shot('20260919-09-audit.png')
                assert errors == [], errors
                browser.close()
            report = {'verified_at': time.strftime('%Y-%m-%d %H:%M:%S %z'), 'platform': platform.platform(),
                      'server_python': subprocess.check_output([interpreter, '--version'], text=True).strip(),
                      'command': 'python main.py', 'isolated_data': True, 'mock': False,
                      'inputs': {'teachers': 7, 'students': 6, 'rooms': 1},
                      'results': {'groups': 1, 'students_assigned': 6, 'status': 'published', 'page_errors': errors},
                      'checks': ['HTTP health', 'frontend served', 'initial admin generated', 'login', 'Excel import x3',
                                 'five-step wizard', 'generate draft', 'publish', 'Excel parse', 'Word parse', 'audit page'],
                      'screenshots': {name: hashlib.sha256((images / name).read_bytes()).hexdigest() for name in screenshots}}
            (evidence / 'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
            print('PASS: python main.py; 3 Excel imports; wizard; 1 published group / 6 students; Word+Excel; 9 screenshots; zero page errors')
        finally:
            server.terminate()
            server.wait(timeout=15)


if __name__ == '__main__':
    main()

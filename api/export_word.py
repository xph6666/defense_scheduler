"""排期结果导出为 Word 时间安排表。

版式对齐学院归档样例（排答辩文档/生成文档/ 下的真实安排表）：
- 标题、副标题居中，正文为两列表格，每组一块（组号/时间/地点/组长|主席/专家/秘书/学生）；
- 结构标签与时间用红色；导师与其学生用同一颜色标注——预答辩/中期时导师在组内
  可直观看出师生配对，正式答辩（导师回避）时学生保留颜色便于人工核对回避是否正确。
"""
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .schedule_integrity import export_groups
from .export_colors import LABEL_RED, build_mentor_color_map

WEEKDAY_LABELS = '一二三四五六日'


def _styled_run(paragraph, text, color=None, bold=False, size=12):
    """写入一个 run 并设置中文字体（宋体需要单独设置 eastAsia 才对中文生效）"""
    run = paragraph.add_run(text)
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def _format_time_label(raw):
    """'2026-04-07 09:00-12:00' → '2026年4月7日 周二 09:00-12:00'"""
    if not raw:
        return '待定'
    try:
        date_part, clock = raw.split(' ', 1)
        day = datetime.strptime(date_part, '%Y-%m-%d').date()
        return f'{day.year}年{day.month}月{day.day}日 周{WEEKDAY_LABELS[day.weekday()]} {clock}'
    except (ValueError, IndexError):
        return raw


def _label_cell(cell, text):
    paragraph = cell.paragraphs[0]
    _styled_run(paragraph, text, color=LABEL_RED, bold=True)


def _names_with_colors(cell, names_with_color, separator='、'):
    """在单元格里逐人写 run，各自带颜色，分隔符保持黑色"""
    paragraph = cell.paragraphs[0]
    for index, (name, color) in enumerate(names_with_color):
        if index:
            _styled_run(paragraph, separator)
        _styled_run(paragraph, name, color=color)


def export_schedule_word(schedule_version, defense_label):
    """生成 Word 时间安排表文档对象；调用方负责写入响应。"""
    groups = list(export_groups(schedule_version))
    color_map = build_mentor_color_map(groups)
    student_total = sum(len(group.students.all()) for group in groups)
    chair_label = '主席' if schedule_version.defense_type == 'formal' else '组长'

    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    suffix = '（草稿，未发布）' if schedule_version.status != 'published' else ''
    _styled_run(title, f'软件工程硕士研究生{defense_label}时间安排{suffix}', bold=True, size=16)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _styled_run(subtitle, f'共 {len(groups)} 组，{student_total} 名学生', size=12)

    note = doc.add_paragraph()
    _styled_run(note, '注：请提前半小时到达答辩教室，准备好论文与 PPT。', size=10.5)

    table = doc.add_table(rows=0, cols=2)
    table.style = 'Table Grid'

    for index, group in enumerate(groups, start=1):
        students = sorted(
            group.students.all(),
            key=lambda s: ((s.mentor_name or ''), s.id),
        )

        header_row = table.add_row()
        _label_cell(header_row.cells[0], f'第{index}组')

        time_row = table.add_row()
        _label_cell(time_row.cells[0], '时间')
        _styled_run(time_row.cells[1].paragraphs[0], _format_time_label(group.time), color=LABEL_RED)

        room_row = table.add_row()
        _label_cell(room_row.cells[0], '地点')
        room_text = f"{group.campus or ''}{group.room.name if group.room else '待定'}"
        _styled_run(room_row.cells[1].paragraphs[0], room_text)

        if group.chair is not None:
            chair_row = table.add_row()
            _label_cell(chair_row.cells[0], chair_label)
            _names_with_colors(
                chair_row.cells[1],
                [(group.chair.name, color_map.get(group.chair.name))],
            )

        expert_row = table.add_row()
        _label_cell(expert_row.cells[0], '专家')
        _names_with_colors(
            expert_row.cells[1],
            [(expert.name, color_map.get(expert.name)) for expert in group.experts.all()],
        )

        secretary_row = table.add_row()
        _label_cell(secretary_row.cells[0], '秘书')
        if group.secretary is not None:
            _styled_run(secretary_row.cells[1].paragraphs[0], group.secretary.name)

        student_row = table.add_row()
        _label_cell(student_row.cells[0], '学生')
        _names_with_colors(
            student_row.cells[1],
            [
                (student.name, color_map.get((student.mentor_name or '').strip()))
                for student in students
            ],
        )

    # 统一列宽：标签列窄、内容列宽（python-docx 需要逐单元格设置才稳定生效）
    for row in table.rows:
        row.cells[0].width = Cm(2.5)
        row.cells[1].width = Cm(13.5)

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _styled_run(footer, '软件学院', size=12)

    footer_date = doc.add_paragraph()
    footer_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    generated = schedule_version.created_at
    _styled_run(footer_date, f'{generated.year}年{generated.month}月{generated.day}日', size=12)

    return doc

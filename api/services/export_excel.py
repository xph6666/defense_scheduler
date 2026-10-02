"""Workbook rendering separated from HTTP and version selection."""
from ..models import Teacher
from ..schedule_integrity import export_groups
from ..mentor_colors import build_mentor_color_map
from scheduling.policies import scenario_policy


def build_schedule_workbook(schedule_version, defense_type, conflicts=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    groups = list(export_groups(schedule_version))
    include_remarks = schedule_version.rules_snapshot.get('include_remarks', True)
    scenario_color = scenario_policy(defense_type)['color'].lstrip('#')
    conflicts = conflicts or []
    error_groups = {g for c in conflicts if c.get('level') == 'error' for g in c.get('relatedGroupIds', [])}
    if any(c.get('level') == 'error' and not c.get('relatedGroupIds') for c in conflicts):
        error_groups = {g.id for g in groups}

    # 导师与其学生同色（与 Word 导出同一套配色语义，色值来自 shared/mentor-colors.json）
    mentor_color_map = build_mentor_color_map(groups)
    mentor_titles = dict(Teacher.objects.values_list('name', 'title'))

    for group in groups:
        sheet = wb.create_sheet(title=f"组{group.group_id}")

        sheet.column_dimensions['A'].width = 15
        sheet.column_dimensions['B'].width = 20
        sheet.column_dimensions['C'].width = 15
        sheet.column_dimensions['D'].width = 25
        if include_remarks:
            sheet.column_dimensions['E'].width = 38

        title_font = Font(bold=True, size=14, color='FFFFFF' if defense_type == 'mid' else '000000')
        title_cell = sheet['A1']
        title_cell.value = f"答辩排期表 - {group.group_id}"
        title_cell.font = title_font
        title_cell.fill = PatternFill(fill_type='solid', fgColor=scenario_color)
        if group.id in error_groups:
            title_cell.font = Font(bold=True, size=14, color='C00000')
        sheet.merge_cells('A1:D1')

        row = 3
        chair_name = group.chair.name if group.chair else '未分配'
        info_data = [
            ('时间', group.time, None),
            ('教室', group.room.name if group.room else '未分配', None),
            ('校区', group.campus, None),
            ('主席/组长', chair_name, mentor_color_map.get(chair_name)),
            ('秘书', group.secretary.name if group.secretary else '未分配', None),
        ]

        for label, value, value_color in info_data:
            sheet[f'A{row}'] = label
            sheet[f'B{row}'] = value
            sheet[f'A{row}'].font = Font(bold=True)
            if value_color:
                sheet[f'B{row}'].font = Font(color=value_color, bold=True)
            row += 1

        sheet[f'A{row}'] = '专家'
        sheet[f'A{row}'].font = Font(bold=True)
        expert_names = [e.name for e in group.experts.all()]
        # 专家逐名着色（富文本）：专家若是某学生导师则与其学生同色
        if expert_names:
            from openpyxl.cell.rich_text import CellRichText, TextBlock
            from openpyxl.cell.text import InlineFont

            rich_parts = []
            for expert_index, expert_name in enumerate(expert_names):
                if expert_index:
                    rich_parts.append('、')
                expert_color = mentor_color_map.get(expert_name)
                if expert_color:
                    rich_parts.append(TextBlock(InlineFont(color=expert_color, b=True), expert_name))
                else:
                    rich_parts.append(expert_name)
            sheet[f'B{row}'] = CellRichText(*rich_parts)
        else:
            sheet[f'B{row}'] = '未分配'
        row += 2

        headers = ['学生姓名', '学生类型', '导师姓名', '导师职称']
        if include_remarks:
            headers.append('备注')
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=row, column=col, value=header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')

        row += 1

        for student in group.students.all():
            mentor_name = (student.mentor_name or '').strip()
            mentor_color = mentor_color_map.get(mentor_name)
            # 只有学生姓名与导师姓名着色（师生配对是唯一靠颜色承载的信息），
            # 学生类型、导师职称等无信息量的列一律黑色，避免整行花花绿绿。
            name_font = Font(color=mentor_color, bold=True) if mentor_color else Font(bold=True)
            plain_font = Font(bold=True)

            cell = sheet.cell(row=row, column=1, value=student.name)
            cell.font = name_font

            cell = sheet.cell(row=row, column=2, value=student.student_type)
            cell.font = plain_font

            cell = sheet.cell(row=row, column=3, value=mentor_name or '未分配')
            cell.font = name_font

            supervisor_title = getattr(student, 'mentor_title', mentor_titles.get(student.mentor_name, ''))
            cell = sheet.cell(row=row, column=4, value=supervisor_title)
            cell.font = plain_font
            if include_remarks:
                sheet.cell(row=row, column=5, value=student.remark or '')

            row += 1

        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for r in range(3, row):
            for c in range(1, 6 if include_remarks else 5):
                cell = sheet.cell(row=r, column=c)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                if group.id in error_groups:
                    cell.font = Font(color='C00000', bold=cell.font.bold)

    roster = wb.create_sheet('教室学生明细')
    headers = ['组别', '时间', '校区', '教室', '学号', '学生', '导师', '秘书']
    if include_remarks:
        headers.append('备注')
    roster.append(headers)
    for group in groups:
        for student in group.students.all():
            values = [group.group_id, group.time, group.campus, group.room.name if group.room else '未分配',
                student.student_no or '', student.name, student.mentor_name,
                group.secretary.name if group.secretary else '未分配']
            if include_remarks:
                values.append(student.remark or '')
            roster.append(values)
            for col in (6, 7):
                roster.cell(roster.max_row, col).font = Font(color=mentor_color_map.get(student.mentor_name, '000000'))
            if group.id in error_groups:
                for cell in roster[roster.max_row]:
                    cell.font = Font(color='C00000')
    if conflicts:
        conflict_sheet = wb.create_sheet('冲突明细')
        conflict_sheet.append(['级别', '组别', '类型', '对象', '说明'])
        for conflict in conflicts:
            conflict_sheet.append([conflict.get('level'), conflict.get('groupName'), conflict.get('type'),
                conflict.get('target'), conflict.get('reason')])
            if conflict.get('level') == 'error':
                for cell in conflict_sheet[conflict_sheet.max_row]:
                    cell.font = Font(color='C00000')
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        if sheet.title in ('教室学生明细', '冲突明细'):
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = PatternFill(fill_type='solid', fgColor=scenario_color)
                cell.font = Font(bold=True, color='FFFFFF' if defense_type == 'mid' else '000000')
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = 24
        for row_cells in sheet:
            for cell in row_cells:
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    cell.data_type = 's'
    return wb

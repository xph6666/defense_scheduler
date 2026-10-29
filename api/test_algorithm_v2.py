"""Executable requirements for the V2 domain and linked resource planning."""
from datetime import datetime
from collections import Counter

from django.test import SimpleTestCase, TestCase

from algorithm import GroupDraft, SchedulingError, TimeRange, assign_teachers, detect_global_conflicts, generate_schedule, parse_room, parse_student, parse_teacher
from scheduling.policies import committee_violations, is_software_teacher, scenario_policy, teacher_can_role


# [TASK-P04 测试] 覆盖专家总人数、角色资格与研究员职称。
class CommitteeV2Tests(SimpleTestCase):
    def rules(self, kind='mid', **overrides):
        return {
            'defense_type': kind, 'start_date': '2026-10-12', 'end_date': '2026-10-12',
            'group_size': 1, 'expert_count': 5,
            'need_chair': True, 'chair_title': '教授' if kind == 'formal' else '副教授',
            'secretary_title': '讲师', 'expert_count_includes_chair': True,
            **overrides,
        }

    def teachers(self, count=9):
        return [
            {'id': i, 'name': f'T{i}', 'title': '教授', 'roles': ['组长', '主席', '专家', '秘书'],
             'is_software_teacher': i <= 5}
            for i in range(1, count + 1)
        ]

    def generate(self, rules=None, teachers=None, students=None):
        return generate_schedule(
            teachers or self.teachers(), students or [{'id': 1, 'name': 'S1'}],
            [{'id': 1, 'name': 'R1', 'capacity': 30}], rules or self.rules(),
        )

    def check(self, group, teachers=None, rules=None, students=None):
        return detect_global_conflicts(
            [group], [parse_student(s) for s in (students or [{'id': 1, 'name': 'S1'}])],
            [parse_teacher(t) for t in (teachers or self.teachers())],
            [parse_room({'id': 1, 'capacity': 30})], rules or self.rules(),
        )

    def group(self, chair=1, experts=None, secretary=6):
        return GroupDraft('G1', None, [1], TimeRange(datetime(2026, 10, 12, 9), datetime(2026, 10, 12, 12)),
                          1, chair, experts if experts is not None else [2, 3, 4], secretary)

    def test_mid_has_four_ordinary_experts_and_one_chair(self):
        result = self.generate()
        self.assertEqual(len(result['groups'][0]['expert_ids']), 4)
        self.assertEqual(result['conflicts'], [])

    def test_personnel_draft_survives_missing_rooms(self):
        result = generate_schedule(self.teachers(), [{'id': 1, 'name': 'S1'}], [], self.rules())
        group = result['groups'][0]
        self.assertEqual(len({group['chair_id'], *group['expert_ids']} - {None}), 5)
        self.assertIsNotNone(group['secretary_id'])
        self.assertIsNone(group['room_id'])
        self.assertFalse(group['time'])
        self.assertEqual({c['type'] for c in result['conflicts']}, {'room_or_time_unavailable'})

    def test_grouping_separates_chair_only_mentors(self):
        teachers = self.teachers(12)
        for teacher in teachers[:2]:
            teacher['roles'] = ['组长']
            teacher['member_eligible'] = False
        students = [{'id': i, 'supervisor_id': 1 if i <= 4 else 2} for i in range(1, 9)]
        result = generate_schedule(teachers, students, [{'id': 1}], self.rules(
            end_date='2026-10-13', group_size=8, group_min=4, group_max=8,
            grouping='supervisor', supervisor_policy='same_group'))
        self.assertEqual(sorted(len(g['student_ids']) for g in result['groups']), [4, 4])
        self.assertEqual(result['conflicts'], [])

    def test_pre_does_not_supply_an_unrelated_chair_to_expert_only_mentors(self):
        teachers = self.teachers(12)
        for teacher in teachers[:8]:
            teacher.update(roles=['普通专家'], expert_only=True)
        students = [{'id': i, 'supervisor_id': i} for i in range(1, 9)]
        result = generate_schedule(teachers, students, [{'id': 1}], self.rules('pre',
            expert_count=4, expert_min=4, group_size=8, group_min=4, group_max=8,
            grouping='supervisor', supervisor_policy='same_group'))
        self.assertTrue(all(g['chair_id'] is None for g in result['groups']))
        self.assertIn('chair_unavailable', {c['type'] for c in result['conflicts']})

    def test_optional_reviewers_leave_future_required_mentors_available(self):
        teachers = self.teachers(14)
        for teacher in teachers:
            teacher['roles'] = (['普通专家'] if teacher['id'] <= 9 else
                                ['组长'] if teacher['id'] <= 12 else ['秘书'])
        teachers[1]['available_time'] = ['2026-10-12 12:00-18:00']
        students = [{'id': 1, 'supervisor_id': 1}, {'id': 2, 'supervisor_id': 2}]
        result = generate_schedule(teachers, students, [{'id': 1}, {'id': 2}], self.rules(
            supervisor_policy='same_group', prefer_full_day_external=False))
        # The second mentor only has the morning; spare them in optional seats.
        for group in result['groups']:
            mentor = group['student_ids'][0]
            self.assertIn(mentor, {group['chair_id'], *group['expert_ids']})

    def test_pre_committee_requires_four_and_can_include_more_required_mentors(self):
        under = committee_violations(self.group(experts=[2, 3]), self.teachers(), self.rules('pre'))
        over = committee_violations(self.group(experts=[2, 3, 4, 5, 7]), self.teachers(), self.rules('pre'))
        self.assertIn('expert_count_out_of_range', {conflict['type'] for conflict in under})
        self.assertNotIn('expert_count_out_of_range', {conflict['type'] for conflict in over})

    def test_generation_accepts_researcher_as_formal_professor_rank(self):
        teachers = self.teachers()
        teachers[0]['title'] = '研究员'
        for teacher in teachers[1:]:
            teacher['title'] = '副教授'
        result = self.generate(self.rules('formal'), teachers=teachers)
        self.assertEqual(result['groups'][0]['chair_id'], 1)
        self.assertEqual(result['conflicts'], [])

    def test_formal_has_exactly_five_reviewers_including_chair(self):
        result = self.generate(self.rules('formal'))
        group = result['groups'][0]
        self.assertEqual(len(set([group['chair_id'], *group['expert_ids']])), 5)
        self.assertEqual(result['conflicts'], [])

    def test_fallback_room_records_its_actual_campus(self):
        students = [{'id': 1, 'name': '跨校区合成学生', 'campus': '兴庆'}]
        rooms = [{'id': 1, 'campus': '创新港', 'name': '异校区合成教室',
                  'available_time': ['2026-10-12 09:00-12:00']}]
        result = generate_schedule(self.teachers(), students, rooms, self.rules())
        self.assertEqual(result['groups'][0]['room_id'], 1)
        self.assertEqual(result['groups'][0]['campus'], '创新港')

    def test_fallback_room_respects_destination_campus_start_date(self):
        students = [{'id': 1, 'name': '跨校区合成学生', 'campus': '兴庆'}]
        rooms = [{'id': 1, 'campus': '创新港', 'name': '异校区合成教室',
                  'available_time': ['2026-10-12 09:00-12:00', '2026-10-13 09:00-12:00']}]
        rules = self.rules(end_date='2026-10-13', campus_start_dates={'创新港': '2026-10-13'})
        result = generate_schedule(self.teachers(), students, rooms, rules)
        self.assertTrue(result['groups'][0]['time'].startswith('2026-10-13'), result)

    def test_legacy_expert_count_keeps_ordinary_expert_semantics(self):
        rules = self.rules(expert_count=2)
        rules.pop('expert_count_includes_chair')
        self.assertEqual(len(self.generate(rules)['groups'][0]['expert_ids']), 2)

    def test_formal_manual_edit_cannot_exceed_five_reviewers(self):
        conflicts = self.check(self.group(experts=[2, 3, 4, 5, 7]), rules=self.rules('formal'))
        self.assertIn('expert_count_out_of_range', {c['type'] for c in conflicts})

    def test_formal_count_validation_follows_configured_target_and_minimum(self):
        rules = self.rules('formal', expert_count=6, expert_min=6)
        valid = self.check(self.group(experts=[2,3,4,5,7], secretary=8), rules=rules)
        self.assertNotIn('expert_count_out_of_range', {c['type'] for c in valid})
        under = self.check(self.group(experts=[2,3,4,5], secretary=8), rules=rules)
        over = self.check(self.group(experts=[2,3,4,5,7,9], secretary=8), rules=rules)
        for conflicts in (under, over):
            self.assertIn('expert_count_out_of_range', {c['type'] for c in conflicts})

    def test_formal_requires_software_college_majority(self):
        teachers = self.teachers()
        for teacher in teachers:
            teacher['is_software_teacher'] = teacher['id'] == 1
        conflicts = self.check(self.group(experts=[2, 3, 4, 5]), teachers, self.rules('formal'))
        self.assertIn('software_majority', {c['type'] for c in conflicts})

    def test_generation_prefers_software_college_reviewers(self):
        teachers = self.teachers()
        for teacher in teachers:
            teacher['is_software_teacher'] = teacher['id'] >= 6
        result = self.generate(self.rules('formal'), teachers)
        group = result['groups'][0]
        self.assertGreaterEqual(len(set([group['chair_id'], *group['expert_ids']]) & {6, 7, 8, 9}), 3)

    def test_same_random_seed_reproduces_group_assignments(self):
        students = [{'id': i, 'name': f'S{i}'} for i in range(1, 9)]
        rules = self.rules('formal', group_size=2, random_seed=938475)

        first = self.generate(rules, students=students)
        second = self.generate(rules, students=students)

        self.assertEqual(first['groups'], second['groups'])
        self.assertEqual(first['random_seed'], 938475)

    def test_inactive_teacher_is_never_automatically_assigned(self):
        teachers = self.teachers()
        teachers[0]['is_active'] = False
        group = self.generate(teachers=teachers)['groups'][0]
        self.assertNotIn(1, [group['chair_id'], *group['expert_ids'], group['secretary_id']])

    def test_chair_only_teacher_is_not_ordinary_member(self):
        teachers = self.teachers()
        teachers[0]['roles'] = ['组长']
        teachers[0]['member_eligible'] = False
        group = self.generate(teachers=teachers)['groups'][0]
        self.assertEqual(group['chair_id'], 1)
        self.assertNotIn(1, group['expert_ids'])

    def test_member_exclusion_is_detected_after_manual_edit(self):
        teachers = self.teachers()
        teachers[1]['member_eligible'] = False
        conflicts = self.check(self.group(), teachers)
        self.assertIn('member_ineligible', {c['type'] for c in conflicts})

    def test_role_qualification_is_detected_after_manual_edit(self):
        teachers = self.teachers()
        teachers[1]['roles'] = ['秘书']
        conflicts = self.check(self.group(), teachers)
        self.assertIn('role_qualification', {c['type'] for c in conflicts})

    def test_bound_secretary_is_reserved_before_member_selection(self):
        students = [{'id': 1, 'name': 'S1', 'secretary_id': 2}]
        group = self.generate(students=students)['groups'][0]
        self.assertEqual(group['secretary_id'], 2)

    def test_formal_preserves_previous_group_label(self):
        students = [{'id': 1, 'name': 'S1', 'previous_group_id': 'G8'},
                    {'id': 2, 'name': 'S2', 'previous_group_id': 'G2'}]
        result = self.generate(self.rules('formal', grouping='secretary'), students=students)
        self.assertEqual({g['student_ids'][0]: g['group_id'] for g in result['groups']}, {1: 'G8', 2: 'G2'})

    def test_shortage_remains_an_explicit_draft_conflict(self):
        result = self.generate(self.rules('formal'), self.teachers(4))
        self.assertIn('expert_count_out_of_range', {c['type'] for c in result['conflicts']})

    def test_mid_student_count_is_checked_with_v2_bounds(self):
        conflicts = self.check(self.group(experts=[2, 3, 4, 5]), rules=self.rules('mid', group_min=10, group_max=13))
        self.assertIn('group_size_out_of_range', {c['type'] for c in conflicts})

    def test_additional_required_mentors_choose_a_large_enough_room(self):
        teachers = self.teachers(9)
        students = [{'id': i, 'name': f'S{i}', 'supervisor_id': i} for i in range(1, 9)]
        rooms = [{'id': 1, 'name': 'Small', 'capacity': 13},
                 {'id': 2, 'name': 'Large', 'capacity': 17}]
        rules = self.rules(group_size=8, group_min=8, group_max=8,
                           grouping='supervisor', supervisor_policy='same_group')
        rules.pop('expert_count_includes_chair')
        result = generate_schedule(teachers, students, rooms, rules)
        # Eight required mentors plus a separate secretary use nine staff seats,
        # even though the ordinary committee target is only four reviewers.
        self.assertEqual(result['groups'][0]['room_id'], 2)
        self.assertEqual(result['conflicts'], [])

    def test_minority_secretary_binding_is_not_hidden_by_majority(self):
        students = [{'id': i, 'name': f'S{i}', 'secretary_id': 6 if i < 3 else 7}
                    for i in range(1, 4)]
        group = self.group()
        group.student_ids = [1, 2, 3]
        conflicts = self.check(group, students=students)
        continuity = [c for c in conflicts if c['type'] == 'secretary_continuity_broken']
        self.assertEqual(len(continuity), 1)
        self.assertEqual(continuity[0]['teacher_name'], 'T7')
        self.assertIn(3, continuity[0]['related_ids'])

    def test_each_distinct_lost_secretary_binding_is_reported(self):
        students = [{'id': i, 'name': f'S{i}', 'secretary_id': i + 6}
                    for i in range(1, 4)]
        group = self.group()
        group.student_ids = [1, 2, 3]
        conflicts = self.check(group, students=students)
        self.assertEqual({c['teacher_name'] for c in conflicts if c['type'] == 'secretary_continuity_broken'},
                         {'T7', 'T8', 'T9'})


class FormalMentorSessionV2Tests(SimpleTestCase):
    def test_cover_expands_chairs_into_expert_seats_when_mentors_require_it(self):
        from scheduling.session_cover import solve_session_cover
        teachers=[parse_teacher({'id':i,'title':'教授','is_software_teacher':True,
                  'roles':['主席'] if i<=4 else ['普通专家'] if i<=12 else ['秘书']}) for i in range(1,15)]
        students=[parse_student({'id':i,'supervisor_id':i}) for i in range(1,5)]
        groups=[GroupDraft('G1',None,[1,2]),GroupDraft('G2',None,[3,4])]
        rooms=[parse_room({'id':i,'capacity':30}) for i in (1,2)]
        slot=TimeRange(datetime(2026,10,12,9),datetime(2026,10,12,12))
        rules={'defense_type':'formal','start_date':'2026-10-12','end_date':'2026-10-12',
               'expert_count':5,'expert_count_includes_chair':True,'need_chair':True,
               'chair_title':'教授','secretary_title':'讲师','avoid_supervisor':True,
               'supervisor_policy':'avoid','formal_mentor_same_session':True,
               'soft_weights':{'balance_software_participation':0}}
        result=solve_session_cover(groups,students,teachers,rooms,rules,[slot],42)
        self.assertIsNotNone(result)
        self.assertEqual(detect_global_conflicts(result,students,teachers,rooms,rules),[])
        self.assertTrue(any(tid<=4 for g in result for tid in g.expert_ids))

    def test_equivalent_sessions_preserve_capacity_and_do_not_share_staff(self):
        self.assert_repeated_session_cover(False)

    def test_equal_teacher_availability_keeps_different_campus_room_windows(self):
        self.assert_repeated_session_cover(True)

    def assert_repeated_session_cover(self, campus_windows):
        from scheduling.session_cover import solve_session_cover
        teachers=[parse_teacher({'id':i,'title':'教授','is_software_teacher':True,
                    'roles':['主席','普通专家'] if i<=12 else ['秘书']}) for i in range(1,15)]
        students=[parse_student({'id':i,'supervisor_id':i,'campus':'A' if i<=2 else 'B'}) for i in range(1,5)]
        groups=[GroupDraft(f'G{s.id}',s.campus if campus_windows else None,[s.id]) for s in students]
        if not campus_windows:
            for student in students:student.campus=None
        slots=[TimeRange(datetime(2026,10,day,9),datetime(2026,10,day,12)) for day in (12,13)]
        if campus_windows:
            rooms=[parse_room({'id':i,'capacity':30,'campus':'A' if i<=2 else 'B',
                              'available_time':[slots[0 if i<=2 else 1].label]}) for i in range(1,5)]
        else:
            rooms=[parse_room({'id':i,'capacity':30}) for i in (1,2)]
        rules={'defense_type':'formal','start_date':'2026-10-12','end_date':'2026-10-13',
               'expert_count':5,'expert_count_includes_chair':True,'need_chair':True,
               'chair_title':'教授','secretary_title':'讲师','avoid_supervisor':True,
               'supervisor_policy':'avoid','formal_mentor_same_session':True,
               'soft_weights':{'balance_software_participation':0}}
        result=solve_session_cover(groups,students,teachers,rooms,rules,slots,42)
        self.assertIsNotNone(result)
        self.assertEqual(len(result),4)
        self.assertEqual(detect_global_conflicts(result,students,teachers,rooms,rules),[])
        self.assertEqual(sorted(sum(g.time_slot==slot for g in result) for slot in slots),[2,2])
        if campus_windows:
            self.assertTrue(all(g.time_slot==slots[0 if g.campus=='A' else 1] for g in result))
        self.assertTrue(all(g.time_slot is None for g in groups))

    def test_exact_cover_schedules_an_odd_cohort_as_one_valid_triple(self):
        from scheduling.session_cover import solve_session_cover
        teachers=[parse_teacher({'id':i,'name':f'CoverT{i}','title':'教授','is_software_teacher':True,
                  'roles':['主席','普通专家'] if i<=15 else ['秘书']}) for i in range(1,19)]
        students=[parse_student({'id':i,'supervisor_id':i}) for i in range(1,4)]
        groups=[GroupDraft(f'G{i}',None,[i]) for i in range(1,4)]
        slot=TimeRange(datetime(2026,10,12,9),datetime(2026,10,12,12))
        rooms=[parse_room({'id':i,'capacity':30}) for i in range(1,4)]
        rules={'defense_type':'formal','start_date':'2026-10-12','end_date':'2026-10-12',
               'expert_count':5,'expert_count_includes_chair':True,'need_chair':True,
               'chair_title':'教授','secretary_title':'讲师','avoid_supervisor':True,
               'supervisor_policy':'avoid','formal_mentor_same_session':True,
               'soft_weights':{'balance_software_participation':0}}
        result=solve_session_cover(groups,students,teachers,rooms,rules,[slot],42)
        self.assertIsNotNone(result)
        self.assertEqual(detect_global_conflicts(result,students,teachers,rooms,rules),[])
        self.assertTrue(all(group.time_slot is None for group in groups))
        self.assertIsNone(solve_session_cover(groups,students,teachers,rooms[:2],rules,[slot],42))

    def test_generation_pairs_complementary_groups_instead_of_trapping_a_shared_mentor(self):
        teachers = [{'id': i, 'name': f'SessionT{i}', 'title': '教授',
                     'is_software_teacher': True,
                     'roles': ['主席', '普通专家'] if i <= 16 else ['秘书']}
                    for i in range(1, 21)]
        students = [{'id': i * 2 + j + 1, 'supervisor_id': tid, 'previous_group_id': f'G{i+1}'}
                    for i, mentors in enumerate(((1, 2), (2, 3), (3, 4), (4, 1)))
                    for j, tid in enumerate(mentors)]
        rules = {'defense_type': 'formal', 'start_date': '2026-10-12', 'end_date': '2026-10-12',
                 'group_size': 2, 'expert_count': 5, 'expert_count_includes_chair': True,
                 'need_chair': True, 'chair_title': '教授', 'secretary_title': '讲师',
                 'supervisor_policy': 'avoid', 'avoid_supervisor': True,
                 'grouping': 'secretary', 'formal_mentor_same_session': True, 'random_seed': 42,
                 'soft_weights': {'balance_software_participation': 0}}
        result = generate_schedule(teachers, students, [{'id': 1}, {'id': 2}], rules)
        self.assertEqual(result['conflicts'], [])
        sessions = {g['group_id']: g['time'] for g in result['groups']}
        self.assertEqual(sessions['G1'], sessions['G3'])
        self.assertEqual(sessions['G2'], sessions['G4'])
        self.assertNotEqual(sessions['G1'], sessions['G2'])

    def test_session_matching_reserves_a_third_group_for_a_shared_external_mentor(self):
        from scheduling.session_assignment import match_formal_session
        teachers = []
        for i in range(1, 19):
            roles = ['主席'] if i in (1, 2) else ['普通专家'] if i < 16 else ['秘书']
            if i == 15:
                roles.append('主席')
            teachers.append(parse_teacher({'id': i, 'title': '教授', 'roles': roles,
                'is_software_teacher': i != 3, 'expert_only': i == 3}))
        students = [parse_student({'id': i, 'supervisor_id': mentor})
                    for i, mentor in enumerate((1, 3, 2, 3, 4), 1)]
        groups = [GroupDraft('G1', None, [1, 2]), GroupDraft('G2', None, [3, 4]),
                  GroupDraft('G3', None, [5])]
        slot = TimeRange(datetime(2026, 10, 12, 9), datetime(2026, 10, 12, 12))
        rules = {'defense_type': 'formal', 'need_chair': True, 'expert_count': 5,
                 'expert_count_includes_chair': True, 'chair_title': '教授',
                 'formal_software_min': 3, 'avoid_supervisor': True}
        result = match_formal_session(groups, students, teachers, rules, slot)
        self.assertIsNotNone(result)
        self.assertIn(3, result[2].expert_ids)
        self.assertNotIn(3, result[0].expert_ids + result[1].expert_ids)
        people = [tid for g in result for tid in [g.chair_id, *g.expert_ids, g.secretary_id]]
        self.assertEqual(len(people), len(set(people)))
        self.assertTrue(all(g.time_slot is None for g in groups))
        self.assertIsNone(match_formal_session(groups[:2], students[:4], teachers, rules, slot))
        self.assertTrue(all(g.chair_id is None for g in groups))

    def check(self, second_hour=9):
        rules = {'defense_type': 'formal', 'expert_count': 0, 'need_chair': False,
                 'supervisor_policy': 'avoid', 'formal_mentor_same_session': True}
        students = [parse_student({'id': 1, 'name': 'S1', 'supervisor_id': 1}),
                    parse_student({'id': 2, 'name': 'S2', 'supervisor_id': 2})]
        teachers = [parse_teacher({'id': i, 'name': f'T{i}'}) for i in (1, 2, 3, 4)]
        groups = [
            GroupDraft('G1', None, [1], TimeRange(datetime(2026, 10, 12, 9), datetime(2026, 10, 12, 12)),
                       1, None, [2], 3),
            GroupDraft('G2', None, [2], TimeRange(datetime(2026, 10, 12, second_hour), datetime(2026, 10, 12, second_hour + 3)),
                       2, None, [1], 4),
        ]
        return detect_global_conflicts(groups, students, teachers, [parse_room({'id': 1}), parse_room({'id': 2})], rules)

    def test_mentor_reviews_another_group_while_own_student_attends(self):
        self.assertNotIn('formal_mentor_session_missing', {c['type'] for c in self.check()})

    def test_mentor_and_student_at_different_sessions_is_a_hard_conflict(self):
        self.assertIn('formal_mentor_session_missing', {c['type'] for c in self.check(14)})

    def test_generation_aligns_two_groups_and_rotates_their_mentors(self):
        teachers = [{'id': i, 'name': f'T{i}', 'title': '教授', 'is_software_teacher': True}
                    for i in range(1, 13)]
        students = [{'id': 1, 'name': 'S1', 'supervisor_id': 1},
                    {'id': 2, 'name': 'S2', 'supervisor_id': 2}]
        rules = {'defense_type': 'formal', 'start_date': '2026-10-12', 'end_date': '2026-10-12',
                 'group_size': 1, 'expert_count': 5, 'expert_count_includes_chair': True,
                 'need_chair': True, 'chair_title': '教授', 'secretary_title': '讲师',
                 'supervisor_policy': 'avoid', 'formal_mentor_same_session': True}
        result = generate_schedule(teachers, students, [{'id': 1}, {'id': 2}], rules)
        self.assertEqual(result['conflicts'], [])
        self.assertEqual(result['groups'][0]['time'], result['groups'][1]['time'])
        self.assertIn(2, [result['groups'][0]['chair_id'], *result['groups'][0]['expert_ids']])
        self.assertIn(1, [result['groups'][1]['chair_id'], *result['groups'][1]['expert_ids']])


class ScenarioPolicyV2Tests(SimpleTestCase):
    def test_catalog_scenarios_do_not_share_mutable_configuration(self):
        first = scenario_policy('mid')
        first['studentCount']['target'] = 99
        self.assertEqual(scenario_policy('中期答辩')['studentCount'], {'target': 12, 'min': 10, 'max': 12})

    def test_chair_in_member_list_counts_once_and_uses_chair_qualification(self):
        teachers = [{'id': 1, 'roles': ['组长'], 'member_eligible': False, 'title': '教授'},
                    *[{'id': i, 'title': '副教授'} for i in (2, 3, 4, 5)],
                    {'id': 6, 'title': '讲师'}]
        conflicts = committee_violations(
            {'group_id': 'G1', 'chair_id': 1, 'expert_ids': [1, 2, 3, 4, 5], 'secretary_id': 6},
            teachers, {'defense_type': 'pre', 'expert_count_includes_chair': True},
        )
        self.assertEqual(conflicts, [])

    def test_software_flag_adds_to_internal_college_inference(self):
        self.assertTrue(is_software_teacher({'college': '软件学院', 'is_software_teacher': False}))
        self.assertFalse(is_software_teacher({'college': '软件学院', 'is_external': True}))
        self.assertTrue(is_software_teacher({'college': '计算机学院', 'is_software_teacher': True}))

    def test_inactive_camel_case_teacher_and_member_flag_are_respected(self):
        self.assertFalse(teacher_can_role({'isActive': False, 'roles': ['主席']}, 'chair'))
        self.assertFalse(teacher_can_role({'memberEligible': False, 'roles': ['专家']}, 'expert'))

    def test_researcher_titles_meet_equivalent_academic_rank_requirements(self):
        conflicts = committee_violations(
            {'group_id': 'G1', 'chair_id': 1, 'secretary_id': 2},
            [{'id': 1, 'name': '研究员', 'title': '研究员', 'roles': ['主席']},
             {'id': 2, 'name': '副研究员', 'title': '副研究员', 'roles': ['秘书']}],
            {'defense_type': 'formal', 'chair_title': '教授', 'secretary_title': '副教授'},
        )
        self.assertNotIn('role_qualification', {conflict['type'] for conflict in conflicts})


# [TASK-P05 测试] 覆盖外院导师1/2组阈值及容量、校区、组人数边界。
class SchedulingPreferencesV2Tests(SimpleTestCase):
    def external_schedule(self, weight=100, same_day_staff_busy=False, student_count=3,
                          group_min=1, group_max=4, group_size=2):
        teachers = [{'id': i, 'name': f'T{i}', 'title': '教授', 'is_external': i == 1,
                     'available_time': ['2026-10-12 14:00-17:00'] if same_day_staff_busy and i > 1 else []}
                    for i in range(1, 7)]
        students = [{'id': i, 'name': f'S{i}', 'supervisor_id': 1} for i in range(1, student_count + 1)]
        rooms = [{'id': 1, 'capacity': 30, 'available_time': [
            '2026-10-12 09:00-12:00', '2026-10-13 09:00-12:00', '2026-10-12 14:00-17:00']}]
        # This fixture isolates external-visit preferences in the mid scenario;
        # the strict pre-defense mentor committee is covered separately.
        rules = {'defense_type': 'mid', 'start_date': '2026-10-12', 'end_date': '2026-10-13',
                 'group_size': group_size, 'group_min': group_min, 'group_max': group_max,
                 'expert_count': 5, 'expert_count_includes_chair': True,
                 'need_chair': True, 'chair_title': '教授', 'secretary_title': '讲师',
                 'supervisor_policy': 'same_group',
                 'soft_weights': {'external_mentor_concentration': weight}}
        return generate_schedule(teachers, students, rooms, rules)

    def test_external_mentor_with_fewer_than_three_students_stays_in_one_group(self):
        result = self.external_schedule(student_count=2)
        self.assertEqual(len(result['groups']), 1)
        self.assertEqual(result['conflicts'], [])

    def test_external_mentor_with_three_students_is_split_into_two_groups(self):
        result = self.external_schedule()
        self.assertEqual(sorted(len(group['student_ids']) for group in result['groups']), [1, 2])
        self.assertNotIn('external_mentor_group_count', {c['type'] for c in result['conflicts']})
        self.assertEqual(result['conflicts'], [])

    def test_external_mentor_distribution_is_checked_within_each_campus(self):
        teachers = [{'id': i, 'name': f'T{i}', 'title': '教授', 'is_external': i == 1}
                    for i in range(1, 7)]
        students = [{'id': i, 'name': f'S{i}', 'supervisor_id': 1,
                     'campus': '兴庆' if i == 4 else '创新港'} for i in range(1, 5)]
        result = generate_schedule(teachers, students,
            [{'id': 1, 'capacity': 30, 'available_time': [
                '2026-10-12 09:00-12:00', '2026-10-12 14:00-17:00', '2026-10-13 09:00-12:00']}],
            {'defense_type': 'mid', 'start_date': '2026-10-12', 'end_date': '2026-10-13',
             'group_size': 2, 'group_min': 1, 'group_max': 4, 'expert_count': 5,
             'expert_count_includes_chair': True, 'need_chair': True, 'chair_title': '教授',
             'secretary_title': '讲师', 'supervisor_policy': 'same_group'})
        # Three students in one campus need two groups, the other campus's one
        # student needs one. Their independent campus buckets must not conflict.
        self.assertNotIn('external_mentor_group_count', {c['type'] for c in result['conflicts']})

    def test_external_mentor_small_cohort_reports_capacity_conflict_instead_of_splitting(self):
        result = self.external_schedule(student_count=2, group_max=1, group_size=1)
        self.assertEqual(len(result['groups']), 1)
        self.assertIn('group_size_out_of_range', {c['type'] for c in result['conflicts']})

    def test_external_mentor_two_groups_below_minimum_size_reports_conflict(self):
        result = self.external_schedule(student_count=4, group_min=3, group_size=3)
        self.assertEqual(len(result['groups']), 2)
        self.assertIn('group_size_out_of_range', {c['type'] for c in result['conflicts']})
        self.assertNotIn('external_mentor_group_count', {c['type'] for c in result['conflicts']})

    def test_external_mentor_prefers_two_feasible_sessions_on_one_day(self):
        result = self.external_schedule()
        self.assertEqual({g['time'].split(' ')[0] for g in result['groups']}, {'2026-10-12'})
        self.assertEqual(result['conflicts'], [])

    def test_disabled_external_preference_keeps_first_feasible_room_slot(self):
        result = self.external_schedule(0)
        self.assertEqual(result['groups'][1]['time'], '2026-10-13 09:00-12:00')

    def test_external_preference_never_overrides_staff_availability(self):
        result = self.external_schedule(same_day_staff_busy=True)
        self.assertEqual(result['groups'][1]['time'], '2026-10-13 09:00-12:00')
        conflict_types = {c['type'] for c in result['conflicts']}
        self.assertIn('external_mentor_concentration', conflict_types)
        self.assertIn('external_half_day_visit', conflict_types)
        # 核心：可用性约束从未被软偏好打破——外院教师仍然被安排在 13 号

    def test_full_day_room_availability_uses_the_other_unblocked_half_day(self):
        for busy_time, expected in (
            ('2026-10-12 10:00-11:00', '2026-10-12 14:00-17:00'),
            ('2026-10-12 15:00-16:00', '2026-10-12 09:00-12:00'),
        ):
            teachers = [
                {'id': 1, 'name': '专家', 'title': '教授', 'roles': ['专家'],
                 'available_time': [busy_time]},
                {'id': 2, 'name': '秘书', 'title': '讲师', 'roles': ['秘书'],
                 'available_time': [busy_time]},
            ]
            room = {'id': 1, 'name': '全天教室', 'available_time': ['2026-10-12 08:00-18:00']}
            with self.subTest(busy_time=busy_time):
                result = generate_schedule(
                    teachers, [{'id': 1, 'name': '学生'}], [room],
                    {'defense_type': 'pre', 'start_date': '2026-10-12', 'end_date': '2026-10-12',
                     'group_size': 1, 'expert_count': 1, 'need_chair': False,
                     'secretary_title': '讲师'},
                )
                self.assertEqual(result['groups'][0]['time'], expected)
                self.assertEqual(result['conflicts'], [])

    def test_formal_cross_campus_preference_survives_reviewer_load_sort(self):
        teachers = [
            parse_teacher({'id': i, 'name': f'T{i}', 'title': '教授',
                           'is_software_teacher': True, 'roles': ['专家']})
            for i in range(1, 9)
        ]
        teachers.append(parse_teacher({'id': 9, 'name': '秘书', 'title': '讲师', 'roles': ['秘书']}))
        slot = TimeRange(datetime(2026, 10, 12, 9), datetime(2026, 10, 12, 12))
        group = GroupDraft('G1', '创新港', [1], slot)
        rules = {
            'defense_type': 'formal', 'expert_count': 3, 'expert_count_includes_chair': True,
            'need_chair': False, 'secretary_title': '讲师',
            'soft_weights': {'avoid_cross_campus': 80},
        }
        teacher_day_campus = {
            (teacher_id, slot.start.date()): {'兴庆'} for teacher_id in (1, 2, 3)
        }
        _, experts, _, _ = assign_teachers(
            group, [parse_student({'id': 1, 'name': '学生'})], teachers, {}, rules,
            teacher_day_campus=teacher_day_campus,
        )
        self.assertFalse({1, 2, 3} & set(experts))

    def software_schedule(self, balance, seed=None):
        teachers = [{'id': i, 'name': f'T{i}', 'title': '教授', 'is_software_teacher': True,
                     'roles': ['主席', '专家']} for i in range(1, 9)]
        teachers.append({'id': 9, 'name': 'Secretary', 'title': '讲师', 'roles': ['秘书']})
        students = [{'id': i, 'name': f'S{i}'} for i in range(1, 4)]
        rules = {'defense_type': 'formal', 'start_date': '2026-10-12', 'end_date': '2026-10-13',
                 'group_size': 1, 'expert_count': 5, 'expert_count_includes_chair': True,
                 'need_chair': True, 'chair_title': '教授', 'secretary_title': '讲师',
                 'soft_weights': {'balance_software_participation': balance}}
        if seed is not None:
            rules['random_seed'] = seed
        return generate_schedule(teachers, students, [{'id': 1}], rules)

    def test_software_reviewers_rotate_with_at_most_one_count_difference(self):
        result = self.software_schedule(100)
        loads = Counter(tid for group in result['groups'] for tid in [group['chair_id'], *group['expert_ids']])
        self.assertLessEqual(max(loads[i] for i in range(1, 9)) - min(loads[i] for i in range(1, 9)), 1)
        self.assertEqual(result['conflicts'], [])

    def test_random_seed_changes_equal_load_reviewer_selection_reproducibly(self):
        first = self.software_schedule(100, seed=21)
        repeated = self.software_schedule(100, seed=21)
        alternate = self.software_schedule(100, seed=22)

        self.assertEqual(first['groups'], repeated['groups'])
        self.assertNotEqual(first['groups'], alternate['groups'])

    def test_disabled_software_balance_keeps_stable_candidate_order(self):
        result = self.software_schedule(0)
        self.assertEqual([group['chair_id'] for group in result['groups']], [1, 1, 1])


class CampusStartDatesV2Tests(SimpleTestCase):
    def generate(self, dates):
        teachers = [{'id': i, 'name': f'T{i}', 'title': '教授'} for i in range(1, 7)]
        students = [{'id': i, 'name': f'S{i}', 'campus': '兴庆' if i <= 10 else '创新港'}
                    for i in range(1, 21)]
        rooms = [{'id': 1, 'campus': '创新港', 'capacity': 30},
                 {'id': 2, 'campus': '兴庆', 'capacity': 30}]
        rules = {'defense_type': 'mid', 'start_date': '2026-10-12', 'end_date': '2026-10-13',
                 'group_size': 12, 'group_min': 10, 'group_max': 13, 'grouping': 'supervisor',
                 'expert_count': 5, 'expert_count_includes_chair': True,
                 'need_chair': True, 'chair_title': '副教授', 'secretary_title': '讲师',
                 'campus_start_dates': dates}
        return generate_schedule(teachers, students, rooms, rules)

    def test_two_campuses_can_start_on_different_dates(self):
        result = self.generate({'创新港': '2026-10-13', '兴庆': '2026-10-12'})
        self.assertEqual({group['campus']: group['time'].split(' ')[0] for group in result['groups']},
                         {'创新港': '2026-10-13', '兴庆': '2026-10-12'})
        self.assertEqual(result['conflicts'], [])

    def test_unspecified_campus_keeps_global_start_date(self):
        result = self.generate({'创新港': '2026-10-13'})
        self.assertEqual({group['campus']: group['time'].split(' ')[0] for group in result['groups']},
                         {'创新港': '2026-10-13', '兴庆': '2026-10-12'})

    def test_invalid_campus_dates_are_rejected_before_scheduling(self):
        for dates in ({'未知校区': '2026-10-12'}, {'兴庆': '2026-02-31'},
                      {'兴庆': '2026-10-11'}, {'兴庆': '2026-10-14'}, ['2026-10-12']):
            with self.subTest(dates=dates), self.assertRaises(SchedulingError):
                self.generate(dates)

    def test_manual_time_before_campus_start_is_reported(self):
        group = GroupDraft('G1', '创新港', [1],
                           TimeRange(datetime(2026, 10, 12, 9), datetime(2026, 10, 12, 12)),
                           1, 1, [2, 3, 4, 5], 6)
        rules = {'defense_type': 'mid', 'start_date': '2026-10-12', 'end_date': '2026-10-13',
                 'expert_count': 5, 'expert_count_includes_chair': True,
                 'need_chair': True, 'chair_title': '副教授', 'secretary_title': '讲师',
                 'campus_start_dates': {'创新港': '2026-10-13'}}
        conflicts = detect_global_conflicts([group], [parse_student({'id': 1, 'campus': '创新港'})],
                    [parse_teacher({'id': i, 'title': '教授'}) for i in range(1, 7)],
                    [parse_room({'id': 1, 'campus': '创新港'})], rules)
        self.assertTrue(any(c['type'] == 'room_or_time_unavailable' and '校区' in c['description'] for c in conflicts))


class LinkedResourcePlanningV2Tests(SimpleTestCase):
    def test_full_reviewer_capacity_can_fail_when_software_seat_capacity_passes(self):
        from scheduling.joint_groups import software_capacity_is_insufficient,reviewer_capacity_is_insufficient
        teachers=[parse_teacher({'id':i,'roles':['主席','普通专家'],'is_software_teacher':True}) for i in range(1,8)]
        students=[parse_student({'id':1,'supervisor_id':1}),parse_student({'id':2,'supervisor_id':2})]
        batches=[[s] for s in students]
        self.assertFalse(software_capacity_is_insufficient(batches,students,teachers))
        self.assertTrue(reviewer_capacity_is_insufficient(batches,students,teachers))

    def test_joint_grouping_preserves_enough_ordinary_mentors_for_formal_review(self):
        self.assert_joint_grouping_preserves_formal_review(False)

    def test_joint_grouping_still_supports_formal_when_chairs_can_be_experts(self):
        self.assert_joint_grouping_preserves_formal_review(True)

    def assert_joint_grouping_preserves_formal_review(self, chair_members):
        from scheduling.joint_groups import software_capacity_is_insufficient
        teachers=[{'id':i,'name':f'JointT{i}','title':'教授',
                   'is_software_teacher':i<=8,'is_external':11<=i<=14,
                   'member_eligible':i>2 or chair_members,
                   'expert_only':11<=i<=14,
                   'roles':['主席','组长'] if i<=2 else ['普通专家'] if i<=8 or i>=11 else ['秘书']}
                  for i in range(1,15)]
        mentors=[1]*4+[2]*4+[tid for tid in range(3,9) for _ in range(2)]
        students=[{'id':i+1,'supervisor_id':tid,'campus':'创新港'} for i,tid in enumerate(mentors)]
        rooms=[{'id':i,'campus':'创新港','capacity':30} for i in (1,2)]
        rules={'defense_type':'pre','start_date':'2026-10-12','end_date':'2026-10-12',
               'group_size':5,'group_min':4,'group_max':8,'expert_count':4,'expert_min':4,'expert_count_includes_chair':True,
               'need_chair':True,'chair_title':'副教授','secretary_title':'讲师',
               'supervisor_policy':'same_group','grouping':'supervisor','reserve_formal_resources':True,
               'refresh_secretary_bindings':True,'policy_version':2,
               'soft_weights':{'balance_student_count':100,'external_mentor_concentration':0},
               'prefer_full_day_external':False}
        pre=generate_schedule(teachers,students,rooms,rules)
        self.assertEqual(pre['conflicts'],[])
        by_id={s['id']:s for s in students}
        batches=[[parse_student(by_id[i]) for i in g['student_ids']] for g in pre['groups']]
        self.assertFalse(software_capacity_is_insufficient(batches,[parse_student(s) for s in students],
                                                        [parse_teacher(t) for t in teachers]))
        formal_students=[{**by_id[i],'secretary_id':g['secretary_id'],'previous_group_id':g['group_id']}
                         for g in pre['groups'] for i in g['student_ids']]
        formal=generate_schedule(teachers,formal_students,rooms,{**rules,'defense_type':'formal',
                'expert_count':5,'expert_min':5,
                'start_date':'2026-10-13','end_date':'2026-10-13','chair_title':'教授',
                'supervisor_policy':'avoid','avoid_supervisor':True,'grouping':'secretary',
                'formal_mentor_same_session':True,'soft_weights':{'balance_software_participation':0,
                                                              'external_mentor_concentration':0}})
        self.assertEqual(formal['conflicts'],[])
        self.assertEqual(Counter(i for g in formal['groups'] for i in g['student_ids']),Counter(by_id.keys()))

    def generate_linked(self, old_secretary=None, external_mentor=False):
        teachers = [{'id': i, 'name': f'T{i:02}', 'title': '教授', 'college': '软件学院',
                     'is_software_teacher': True, 'roles': ['主席', '组长', '普通专家', '秘书']}
                    for i in range(1, 15)]
        if external_mentor:
            teachers[1].update(title='讲师', college='外院', is_external=True, is_software_teacher=False)
        mentors = [1,1,2,2,3,4,5,6,7,8]
        students = [{'id': i, 'name': f'S{i:02}', 'campus': '创新港',
                     'supervisor_id': mentors[i-1],
                     'secretary_id': old_secretary}
                    for i in range(1, 11)]
        rooms = [{'id': i, 'campus': '创新港', 'capacity': 30} for i in (1, 2)]
        rules = {'defense_type': 'pre', 'start_date': '2026-10-12', 'end_date': '2026-10-12',
                 'group_size': 6, 'group_min': 4, 'group_max': 8, 'expert_count': 4, 'expert_min': 4,
                 'expert_count_includes_chair': True, 'need_chair': True, 'chair_title': '副教授',
                 'secretary_title': '讲师', 'supervisor_policy': 'same_group', 'grouping': 'supervisor',
                 'prefer_senior': True, 'policy_version': 2,
                 'reserve_formal_resources': True, 'refresh_secretary_bindings': True}
        pre = generate_schedule(teachers, students, rooms, rules)
        formal_students = []
        for group in pre['groups']:
            formal_students.extend({**student, 'secretary_id': group['secretary_id'],
                                    'previous_group_id': group['group_id']}
                                   for student in students if student['id'] in group['student_ids'])
        formal = generate_schedule(teachers, formal_students, rooms, {
            **rules, 'defense_type': 'formal', 'start_date': '2026-10-13', 'end_date': '2026-10-13',
            'expert_count': 5, 'expert_min': 5, 'chair_title': '教授', 'supervisor_policy': 'avoid',
            'grouping': 'secretary', 'formal_mentor_same_session': True})
        return pre, formal

    def assert_valid_linked(self, pre, formal):
        self.assertEqual(pre['conflicts'], [])
        self.assertEqual(formal['conflicts'], [])
        self.assertGreaterEqual(len({g['secretary_id'] for g in pre['groups']}), 2)
        self.assertFalse({1, 2} & {g['secretary_id'] for g in pre['groups']})
        self.assertEqual(len({g['time'] for g in formal['groups']}), 1)
        self.assertEqual(len({g['room_id'] for g in formal['groups']}), 2)
        self.assertEqual({g['group_id']: g['student_ids'] for g in pre['groups']},
                         {g['group_id']: g['student_ids'] for g in formal['groups']})

    def test_ten_students_keep_secretaries_and_mentors_in_one_formal_session(self):
        self.assert_valid_linked(*self.generate_linked())

    def test_explicit_replanning_refreshes_a_shared_old_secretary(self):
        self.assert_valid_linked(*self.generate_linked(old_secretary=14))

    def test_less_senior_external_mentor_still_gets_a_formal_review_seat(self):
        pre, formal = self.generate_linked(external_mentor=True)
        self.assertEqual(pre['conflicts'], [])
        external_student_groups = [group for group in pre['groups']
                                   if 3 in group['student_ids']]
        self.assertEqual(len(external_student_groups), 1)
        external_review_groups = [group for group in formal['groups']
                                  if group['chair_id'] == 2 or 2 in group['expert_ids']]
        self.assertTrue(external_review_groups)
        self.assertNotIn('formal_mentor_session_missing', {conflict['type'] for conflict in formal['conflicts']})


class LinkedResourceServiceV2Tests(TestCase):
    def test_no_rooms_keeps_complete_personnel_but_blocks_publication_and_export(self):
        from django.contrib.auth.models import User
        from rest_framework.test import APIClient
        from .models import Student, Teacher
        client = APIClient()
        client.force_authenticate(User.objects.create_user('personnel-draft', is_staff=True))
        teachers = [Teacher.objects.create(name=f'DraftT{i}', title='教授', college='软件学院',
                    roles=['组长', '普通专家', '秘书']) for i in range(7)]
        for i in range(6):
            Student.objects.create(name=f'DraftS{i}', student_no=f'DRAFT{i}', campus='创新港',
                                   mentor=teachers[i % 4], defense_types=['预答辩'])
        response = client.post('/api/schedule/generate/', {
            'rules': {'defense_type': 'pre', 'policy_version': 2,
                      'start_date': '2026-10-12', 'end_date': '2026-10-12'}}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        group = response.data['groups'][0]
        self.assertEqual(len(group['students']), 6)
        self.assertEqual(len({group['chairId'], *[t['id'] for t in group['teachers']]} - {None}), 4)
        self.assertIsNotNone(group['secretaryId'])
        self.assertFalse(group['date'])
        version_id = response.data['versionId']
        published = client.post('/api/schedule/publish/', {
            'version_id': version_id, 'expected_revision': response.data['revision']}, format='json')
        self.assertEqual(published.status_code, 400)
        for endpoint in ('export-excel', 'export_word'):
            exported = client.get(f'/api/schedule/{endpoint}/', {'version_id': version_id, 'defense_type': 'pre'})
            self.assertEqual(exported.status_code, 400)

    def test_linked_replan_after_published_shared_secretary_is_valid_and_keeps_history(self):
        from django.contrib.auth.models import User
        from rest_framework.test import APIClient
        from .models import Group, Room, ScheduleVersion, Student, Teacher
        from .schedule_integrity import export_groups

        client = APIClient()
        client.force_authenticate(User.objects.create_user('resource-planner', is_staff=True))
        teachers = [Teacher.objects.create(name=f'T{i:02}', title='教授', college='软件学院',
                    roles=['主席', '组长', '普通专家', '秘书']) for i in range(1, 15)]
        students = [Student.objects.create(name=f'S{i:02}', student_no=f'LINKED{i}', campus='创新港',
                    mentor=teachers[[0,0,1,1,2,3,4,5,6,7][i-1]], bound_secretary=teachers[-1],
                    defense_types=['预答辩', '正式答辩']) for i in range(1, 11)]
        rooms = [Room.objects.create(name=f'R{i}', campus='创新港', capacity=30) for i in (1, 2)]
        pre_rules = {'defense_type': 'pre', 'policy_version': 2,
                     'start_date': '2026-10-12', 'end_date': '2026-10-12'}
        from .services.rules import normalize_schedule_rules
        old = ScheduleVersion.objects.create(defense_type='pre', rules_snapshot=normalize_schedule_rules(pre_rules))
        for index, time in enumerate(('09:00-12:00', '14:00-17:00')):
            group = Group.objects.create(schedule_version=old, group_id=f'G{index + 1}',
                        time=f'2026-10-12 {time}', campus='创新港', room=rooms[0],
                        chair=teachers[index*4], secretary=teachers[-1])
            group.experts.set(teachers[index*4+1:index*4+4])
            group.students.set(students[index * 6:(index + 1) * 6])
        published = client.post('/api/schedule/publish/',
                               {'version_id': old.pk, 'expected_revision': 0}, format='json')
        self.assertEqual(published.status_code, 200, published.data)
        response = client.post('/api/schedule/generate-linked/', {
            'pre_rules': pre_rules,
            'formal_rules': {'defense_type': 'formal', 'policy_version': 2,
                            'start_date': '2026-10-13', 'end_date': '2026-10-13'},
            'request_key': 'shared-secretary-replan'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        for stage in ('pre', 'formal'):
            self.assertFalse([c for c in response.data[stage]['conflicts'] if c['level'] == 'error'], response.data)
        self.assertEqual(len({g['secretaryId'] for g in response.data['pre']['groups']}), 2)
        self.assertEqual(len({(g['date'], g['timeRange']) for g in response.data['formal']['groups']}), 1)
        old.refresh_from_db()
        self.assertEqual(old.status, 'published')
        self.assertFalse(old.is_current)
        self.assertEqual({s.bound_secretary_id for g in export_groups(old) for s in g.students.all()}, {teachers[-1].pk})


# [TASK-P09/P12 测试] 覆盖版本评分数据、Word备注/冲突标记及关系导出。
class FormalSecretaryExportV2Tests(TestCase):
    def setUp(self):
        from .models import Group, ScheduleVersion, Student, Teacher
        from .schedule_integrity import capture_export_groups

        self.old_secretary = Teacher.objects.create(name='预答辩原秘书', title='讲师', roles=['秘书'])
        self.formal_secretary = Teacher.objects.create(name='正式答辩改选秘书', title='讲师', roles=['秘书'])
        self.later_secretary = Teacher.objects.create(name='下一轮预答辩秘书', title='讲师', roles=['秘书'])
        self.student = Student.objects.create(name='版本秘书学生', student_no='FORMAL-SECRETARY',
            campus='创新港', defense_types=['预答辩', '正式答辩'],
            bound_secretary=self.old_secretary, secretary_name=self.old_secretary.name)
        self.source = ScheduleVersion.objects.create(defense_type='pre', is_current=False)
        source_group = Group.objects.create(schedule_version=self.source, group_id='G1',
            campus='创新港', time='2026-10-12 09:00-12:00', secretary=self.old_secretary)
        source_group.students.add(self.student)
        self.source.export_snapshot = capture_export_groups(self.source)
        self.source.save(update_fields=['export_snapshot'])
        self.version = ScheduleVersion.objects.create(defense_type='formal', source_pre_version=self.source,
            input_snapshot={'students': [{'id': self.student.pk,
                'secretary_id': self.formal_secretary.pk, 'previous_group_id': 'G1'}]})
        self.group = Group.objects.create(schedule_version=self.version, group_id='G1',
            campus='创新港', time='2026-10-13 09:00-12:00', secretary=self.formal_secretary)
        self.group.students.add(self.student)

    def presentation_student(self):
        from unittest.mock import patch
        from .schedule_views import ScheduleViewSet

        view = ScheduleViewSet()
        with patch.object(view, '_check_conflicts', return_value=[]):
            return view._version_result(self.version)['groups'][0]['students'][0]

    def test_schedule_results_expose_external_mentor_metadata_for_live_and_frozen_versions(self):
        from .models import Teacher
        from .schedule_views import ScheduleViewSet

        mentor = Teacher.objects.create(name='外院导师元数据', title='研究员', is_external=True)
        self.student.mentor = mentor
        self.student.mentor_name = mentor.name
        self.student.save(update_fields=['mentor', 'mentor_name'])
        self.assertIs(self.presentation_student()['mentorIsExternal'], True)

        self.version.input_snapshot = {'teachers': [{'id': mentor.pk, 'is_external': True}]}
        self.version.rules_snapshot = {'include_remarks': False}
        self.version.result_snapshot = {'groups': [{'students': [{'id': self.student.pk, 'mentorId': mentor.pk}]}]}
        self.version.save(update_fields=['input_snapshot', 'rules_snapshot', 'result_snapshot'])
        frozen = ScheduleViewSet()._version_result(self.version)
        self.assertIs(frozen['groups'][0]['students'][0]['mentorIsExternal'], True)
        self.assertFalse(frozen['rules']['include_remarks'])

    def assert_student_binding(self, secretary_id, secretary_name):
        from .schedule_integrity import capture_export_groups, export_groups

        student = export_groups(self.version)[0].students.all()[0]
        self.assertEqual(student.bound_secretary_id, secretary_id)
        self.assertEqual(student.secretary_name, secretary_name)
        captured = capture_export_groups(self.version)[0]['students'][0]
        self.assertEqual(captured['bound_secretary_id'], secretary_id)
        self.assertEqual(captured['secretary_name'], secretary_name)
        displayed = self.presentation_student()
        self.assertEqual(displayed['secretaryId'], secretary_id)
        self.assertEqual(displayed['secretaryName'], secretary_name)

    def test_word_exports_render_arrangement_and_conclusion_forms(self):
        """两种归档 Word：安排表按样例只列姓名，结论表把姓名与导师写在下划线空位上。"""
        from docx.shared import RGBColor
        from .export_conclusion import export_conclusion_document
        from .export_word import export_arrangement_document

        self.student.mentor = self.old_secretary
        self.student.mentor_name = self.old_secretary.name
        self.student.remark = '测试备注'
        self.student.student_no = 'WORD-001'
        self.student.save(update_fields=['mentor', 'mentor_name', 'remark', 'student_no'])
        conflict = {'relatedGroupIds': [self.group.pk], 'level': 'error',
                    'type': '导师冲突', 'reason': '冲突测试'}

        arrangement = export_arrangement_document(self.version, 'formal', True, [conflict])
        table_text = '\n'.join(cell.text for row in arrangement.tables[0].rows for cell in row.cells)
        self.assertIn('版本秘书学生', table_text)
        self.assertIn('（测试备注）', table_text)
        self.assertIn('正式答辩改选秘书', table_text)
        # 学院样例的时间安排表只写姓名和备注，学号与逐人导师关系留给结论表
        self.assertNotIn('WORD-001', table_text)
        self.assertNotIn('导师：', table_text)
        appendix = '\n'.join(paragraph.text for paragraph in arrangement.paragraphs)
        self.assertIn('冲突测试', appendix)
        conflict_run = next(run for paragraph in arrangement.paragraphs
                            for run in paragraph.runs if '冲突测试' in run.text)
        self.assertEqual(conflict_run.font.color.rgb, RGBColor.from_string('C00000'))

        without_remarks = export_arrangement_document(self.version, 'formal', False, [conflict])
        text_without_remarks = '\n'.join(cell.text for row in without_remarks.tables[0].rows
                                         for cell in row.cells)
        self.assertNotIn('测试备注', text_without_remarks)

        conclusion = export_conclusion_document(self.version, 'formal', True)
        lines = [paragraph.text for row in conclusion.tables[0].rows
                 for cell in row.cells for paragraph in cell.paragraphs]
        student_lines = [line for line in lines if line.startswith('学生1')]
        self.assertTrue(student_lines, lines)
        self.assertIn('版本秘书学生', student_lines[0])
        self.assertIn('指导教师', student_lines[0])
        self.assertIn('预答辩原秘书', student_lines[0])

    def test_conclusion_export_pads_blank_student_lines_and_page_breaks_each_group(self):
        from .export_conclusion import MIN_STUDENT_ROWS, export_conclusion_document

        self.group.experts.clear()
        document = export_conclusion_document(self.version, 'formal', True)

        cell = document.tables[0].rows[1].cells[0]
        student_lines = [paragraph.text for paragraph in cell.paragraphs
                         if paragraph.text.startswith('学生')]
        self.assertEqual(len(student_lines), MIN_STUDENT_ROWS)
        self.assertIn('指导教师', student_lines[-1])
        self.assertIn('待定', document.tables[0].rows[0].cells[3].text)

    def test_excel_export_prints_all_columns_together_and_repeats_table_headers(self):
        from .services.export_excel import build_schedule_workbook

        workbook = build_schedule_workbook(self.version, 'formal')

        for sheet in workbook:
            self.assertEqual(sheet.page_setup.orientation, 'landscape')
            self.assertEqual(int(sheet.page_setup.paperSize), int(sheet.PAPERSIZE_A4))
            self.assertEqual(sheet.page_setup.fitToWidth, 1)
            self.assertEqual(sheet.page_setup.fitToHeight, 0)
            self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
            expected_area = f'$A$1:${sheet.cell(1, sheet.max_column).column_letter}${sheet.max_row}'
            self.assertIn(expected_area, str(sheet.print_area))
        group_sheet = workbook['组G1']
        header_row = next(row[0].row for row in group_sheet.iter_rows()
                          if row[0].value == '学生姓名')
        self.assertEqual(group_sheet.print_title_rows, f'${header_row}:${header_row}')
        self.assertEqual(workbook['教室学生明细'].print_title_rows, '$1:$1')

    def test_explicit_formal_binding_is_used_without_writing_global_student(self):
        self.assert_student_binding(self.formal_secretary.pk, self.formal_secretary.name)
        self.student.refresh_from_db()
        self.assertEqual(self.student.bound_secretary_id, self.old_secretary.pk)
        self.assertEqual(self.student.secretary_name, self.old_secretary.name)

    def test_later_pre_rebinding_does_not_change_formal_or_frozen_source_view(self):
        from .models import Student
        from .schedule_integrity import export_groups

        Student.objects.filter(pk=self.student.pk).update(bound_secretary=self.later_secretary,
            secretary_name=self.later_secretary.name)
        self.assert_student_binding(self.formal_secretary.pk, self.formal_secretary.name)
        source_student = export_groups(self.source)[0].students.all()[0]
        self.assertEqual(source_student.bound_secretary_id, self.old_secretary.pk)
        self.assertEqual(source_student.secretary_name, self.old_secretary.name)

    def test_explicit_empty_formal_binding_does_not_fall_back_to_global_secretary(self):
        self.version.input_snapshot['students'][0]['secretary_id'] = None
        self.assert_student_binding(None, '')

    def test_missing_snapshot_teacher_keeps_invalid_id_visible_to_validation(self):
        from .schedule_views import ScheduleViewSet

        missing_id = self.later_secretary.pk + 1000
        self.version.input_snapshot['students'][0]['secretary_id'] = missing_id
        self.assert_student_binding(missing_id, '')
        conflicts = ScheduleViewSet()._check_conflicts(self.version)
        self.assertTrue(any(c['type'] == 'secretary_not_found' for c in conflicts), conflicts)

    def test_frozen_formal_export_does_not_reproject_changed_bindings(self):
        from .models import Student
        from .schedule_integrity import capture_export_groups

        self.version.export_snapshot = capture_export_groups(self.version)
        self.version.status = 'published'
        self.version.save(update_fields=['export_snapshot', 'status'])
        self.version.input_snapshot['students'][0]['secretary_id'] = None
        Student.objects.filter(pk=self.student.pk).update(bound_secretary=self.later_secretary,
            secretary_name=self.later_secretary.name)
        self.assert_student_binding(self.formal_secretary.pk, self.formal_secretary.name)

    def test_legacy_formal_snapshot_without_binding_keeps_existing_student_view(self):
        del self.version.input_snapshot['students'][0]['secretary_id']
        self.assert_student_binding(self.old_secretary.pk, self.old_secretary.name)

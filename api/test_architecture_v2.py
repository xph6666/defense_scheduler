"""Cross-layer regression cases for the V2 upgrade and legacy compatibility."""
import json
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase, SimpleTestCase
from rest_framework.test import APIClient

from .models import Group, OperationLog, Room, ScheduleVersion, Student, Teacher
from .schedule_integrity import capture_export_groups
from .services.rules import normalize_schedule_rules
from .services.scheduling_input import build_algorithm_input, expand_busy_half_days


class SchedulePolicyBoundaryTests(SimpleTestCase):
    def test_api_cannot_lower_v2_hard_constraints(self):
        from rest_framework.exceptions import ValidationError
        for overrides in (
            {'chair_title': ''}, {'secretary_title': '助教'},
            {'course_half_day_blocking': False},
            {'defense_type': 'formal', 'formal_mentor_same_session': False},
            {'defense_type': 'formal', 'expert_count': 6},
            {'defense_type': 'mid', 'group_min': 1, 'group_size': 1, 'group_max': 99},
        ):
            with self.subTest(overrides=overrides), self.assertRaises(ValidationError):
                normalize_schedule_rules({'policy_version': 2, **overrides})

    def test_server_source_reference_cannot_be_supplied_by_client(self):
        internal = {'_source_pre_version_id': 999, 'reserve_formal_resources': True, 'refresh_secretary_bindings': True}
        normalized = normalize_schedule_rules({'policy_version': 2, **internal})
        self.assertTrue(all(key not in normalized for key in internal))

    def test_course_blocks_whole_half_day_and_preserves_day_boundaries(self):
        from algorithm import parse_time_range
        values = expand_busy_half_days(['2026-10-12 10:00-11:00', '2026-10-12 16:00-17:00'])
        ranges = [parse_time_range(v) for v in values]
        self.assertEqual([(r.start.hour, r.end.hour) for r in ranges], [(0, 12), (12, 0)])
        self.assertEqual(ranges[1].end.day, 13)


class ArchitectureV2IntegrationTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(User.objects.create_user('v2admin', is_staff=True))
        self.teachers = [Teacher.objects.create(name=f'T{i:02}', title='教授', college='软件学院',
            roles=['主席', '组长', '普通专家', '秘书']) for i in range(1, 15)]
        self.students = []
        for i in range(12):
            mentor = self.teachers[0 if i < 6 else 1]
            self.students.append(Student.objects.create(name=f'S{i:02}', student_no=f'V2{i}',
                mentor=mentor, mentor_name=mentor.name, campus='创新港', defense_types=['预答辩', '正式答辩']))
        self.rooms = [Room.objects.create(name=f'R{i}', campus='创新港') for i in (1, 2)]

    def rules(self, stage):
        day = '2026-10-12' if stage == 'pre' else '2026-10-19'
        return {'defense_type': stage, 'policy_version': 2, 'start_date': day, 'end_date': day}

    def linked(self, key='linked-v2'):
        return self.client.post('/api/schedule/generate-linked/', {
            'pre_rules': self.rules('pre'), 'formal_rules': self.rules('formal'), 'request_key': key}, format='json')

    def test_linked_generation_pins_source_groups_and_retries_without_new_versions(self):
        previous = self.client.post('/api/schedule/generate/', {'rules': self.rules('pre')}, format='json')
        self.assertEqual(previous.status_code, 200, previous.data)
        first = self.linked()
        self.assertEqual(first.status_code, 200, first.data)
        pre, formal = first.data['pre'], first.data['formal']
        self.assertFalse([c for c in pre['conflicts'] if c['level'] == 'error'], pre['conflicts'])
        self.assertFalse([c for c in formal['conflicts'] if c['level'] == 'error'], formal['conflicts'])
        self.assertEqual(formal['sourcePreVersionId'], pre['versionId'])
        memberships = lambda result: {g['groupName']: sorted(s['id'] for s in g['students']) for g in result['groups']}
        self.assertEqual(memberships(pre), memberships(formal))
        for group in formal['groups']:
            self.assertEqual(len({group['chairId'], *(t['id'] for t in group['teachers'])}), 5)
        second = self.linked()
        self.assertEqual(second.status_code, 200, second.data)
        self.assertEqual(second.data['formal']['versionId'], formal['versionId'])
        self.assertEqual(ScheduleVersion.objects.count(), 3)

    def test_linked_failure_rolls_back_both_versions_and_student_bindings(self):
        Student.objects.update(defense_types=['预答辩'])
        failed = self.linked()
        self.assertEqual(failed.status_code, 400, failed.data)
        self.assertEqual(ScheduleVersion.objects.count(), 0)
        self.assertFalse(Student.objects.exclude(secretary_name='').exists())
        self.assertFalse(Student.objects.exclude(bound_secretary=None).exists())
        self.assertFalse(OperationLog.objects.filter(authoritative=True).exists())

    def test_formal_source_does_not_follow_new_current_pre_version(self):
        initial = self.linked()
        self.assertEqual(initial.status_code, 200, initial.data)
        original = ScheduleVersion.objects.get(pk=initial.data['pre']['versionId'])
        new = self.client.post('/api/schedule/generate/', {'rules': self.rules('pre')}, format='json')
        self.assertEqual(new.status_code, 200, new.data)
        payload, _ = build_algorithm_input('formal', rules={'_source_pre_version_id': original.pk})
        source = {s.pk: g.group_id for g in original.groups.prefetch_related('students') for s in g.students.all()}
        self.assertEqual({s['id']: s['previous_group_id'] for s in payload['students']}, source)
        no_source, _ = build_algorithm_input('formal', rules={'_source_pre_version_id': None})
        self.assertTrue(all(s['previous_group_id'] is None for s in no_source['students']))

    def test_adapter_preserves_explicit_identity_and_role_fields(self):
        student = self.students[0]
        student.mentor_name = '过时的姓名'
        student.save(update_fields=['mentor_name'])
        self.teachers[0].roles = ['组长']
        self.teachers[0].member_eligible = False
        self.teachers[0].save()
        self.teachers[2].is_external = True
        self.teachers[2].save()
        payload, _ = build_algorithm_input('pre')
        self.assertEqual(next(s for s in payload['students'] if s['id'] == student.pk)['supervisor_id'], student.mentor_id)
        self.assertEqual(payload['teachers'][0]['roles'], ['组长'])
        self.assertFalse(payload['teachers'][0]['member_eligible'])
        self.assertFalse(next(t for t in payload['teachers'] if t['id'] == self.teachers[2].pk)['is_software_teacher'])

    def test_new_linked_round_ignores_retired_or_unknown_old_secretaries(self):
        retired = self.teachers[-1]
        retired.is_active = False
        retired.save(update_fields=['is_active'])
        Student.objects.filter(pk=self.students[0].pk).update(bound_secretary=retired, secretary_name=retired.name)
        Student.objects.filter(pk=self.students[1].pk).update(secretary_name='不存在的旧秘书')
        payload, notices = build_algorithm_input('pre', rules={'refresh_secretary_bindings': True, 'policy_version': 2})
        self.assertTrue(all(s['secretary_id'] is None for s in payload['students']))
        self.assertFalse(any(n['type'].startswith('secretary_') for n in notices), notices)
        # A normal standalone pre still preserves and reports old bindings.
        _, standalone_notices = build_algorithm_input('pre', rules={'policy_version': 2})
        self.assertEqual({n['type'] for n in standalone_notices},
            {'secretary_not_found', 'secretary_unavailable_for_defense'})

    def test_published_fk_names_remain_consistent_in_results_and_exports(self):
        generated = self.client.post('/api/schedule/generate/', {'rules': self.rules('pre')}, format='json')
        self.assertEqual(generated.status_code, 200, generated.data)
        student = self.students[0]
        Student.objects.filter(pk=student.pk).update(mentor_name='过期字段')
        published = self.client.post('/api/schedule/publish/', {
            'version_id': generated.data['versionId'], 'expected_revision': 0}, format='json')
        self.assertEqual(published.status_code, 200, published.data)
        version = ScheduleVersion.objects.get(pk=generated.data['versionId'])
        rows = [s for g in published.data['groups'] for s in g['students'] if s['id'] == student.pk]
        self.assertEqual(rows[0]['mentorName'], self.teachers[0].name)
        snapshots = [s for g in version.export_snapshot for s in g['students'] if s['id'] == student.pk]
        self.assertEqual(snapshots[0]['mentor_name'], rows[0]['mentorName'])

    def test_adjust_group_rejects_mismatched_teacher_id_without_mutation(self):
        generated = self.client.post('/api/schedule/generate/', {'rules': self.rules('pre')}, format='json')
        self.assertEqual(generated.status_code, 200, generated.data)
        version = ScheduleVersion.objects.get(pk=generated.data['versionId'])
        group = version.groups.first()
        previous = group.chair_id
        response = self.client.post('/api/schedule/adjust-group/', {'group_id': group.pk,
            'group_data': {'chairId': self.teachers[-1].pk, 'chairman': '错误的姓名'}, 'expected_revision': 0}, format='json')
        self.assertEqual(response.status_code, 400, response.data)
        group.refresh_from_db()
        self.assertEqual(group.chair_id, previous)

    def test_teacher_import_keeps_v2_qualification_columns(self):
        from .tests import make_xlsx_upload
        upload = make_xlsx_upload('教师资格.xlsx', [
            ['姓名', '所属学院', '职称', '可担任角色', '是否启用', '是否适合当组员', '是否软件学院导师'],
            ['导入的主席', '计算机学院', '教授', '主席', '否', '否', '是'],
        ])
        imported = self.client.post('/api/teachers/import_data/', {'file': upload}, format='multipart')
        self.assertEqual(imported.status_code, 200, imported.data)
        teacher = Teacher.objects.get(name='导入的主席')
        self.assertFalse(teacher.is_active)
        self.assertFalse(teacher.member_eligible)
        self.assertTrue(teacher.is_software_teacher)

    def test_pre_secretary_id_edit_updates_only_that_group_binding(self):
        generated = self.client.post('/api/schedule/generate/', {'rules': self.rules('pre')}, format='json')
        self.assertEqual(generated.status_code, 200, generated.data)
        version = ScheduleVersion.objects.get(pk=generated.data['versionId'])
        source, other = list(version.groups.all())
        untouched = dict(other.students.values_list('pk', 'bound_secretary_id'))
        replacement = next(t for t in self.teachers if t.pk not in {source.secretary_id, source.chair_id})
        changed = self.client.post('/api/schedule/adjust-group/', {'group_id': source.pk,
            'group_data': {'secretaryId': replacement.pk}, 'expected_revision': 0}, format='json')
        self.assertEqual(changed.status_code, 200, changed.data)
        self.assertEqual(set(source.students.values_list('bound_secretary_id', flat=True)), {replacement.pk})
        self.assertEqual(dict(other.students.values_list('pk', 'bound_secretary_id')), untouched)

    def test_move_keeps_secretary_and_moves_mentor_then_rechecks_conflicts(self):
        generated = self.client.post('/api/schedule/generate/', {'rules': self.rules('pre')}, format='json')
        self.assertEqual(generated.status_code, 200, generated.data)
        version = ScheduleVersion.objects.get(pk=generated.data['versionId'])
        source, target = list(version.groups.all())
        student = source.students.first()
        old_secretary = source.secretary_id
        different_secretary = next(t for t in self.teachers if t.pk != old_secretary)
        target.secretary = different_secretary
        target.save(update_fields=['secretary'])
        target.students.update(bound_secretary=different_secretary, secretary_name=different_secretary.name)
        target_secretary = different_secretary.pk
        response = self.client.post('/api/schedule/adjust/', {'action': 'move_student', 'student_id': student.pk,
            'from_group_id': source.pk, 'to_group_id': target.pk, 'expected_revision': 0}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        target.refresh_from_db()
        student.refresh_from_db()
        self.assertEqual(student.bound_secretary_id, old_secretary)
        self.assertEqual(target.secretary_id, target_secretary)
        self.assertTrue(target.experts.filter(pk=student.mentor_id).exists() or target.chair_id == student.mentor_id)
        self.assertTrue(response.data['movement']['mentorMoved'])
        self.assertTrue(OperationLog.objects.filter(operator='v2admin', authoritative=True).exists())
        checked = self.client.post('/api/schedule/check-conflicts/', {'defense_type': 'pre'}, format='json')
        current = self.client.get('/api/schedule/current/', {'defense_type': 'pre'})
        self.assertEqual(current.data['conflicts'], checked.data)
        self.assertTrue(any(c['level'] == 'error' and c['type'] == '秘书连续性提示' for c in current.data['conflicts']), current.data['conflicts'])
        # An unrelated edit must not silently erase the preserved relationship.
        changed = self.client.post('/api/schedule/adjust/', {'action': 'change_time', 'group_id': target.pk,
            'new_time': '2026-10-12 14:00-17:00', 'expected_revision': 1}, format='json')
        self.assertEqual(changed.status_code, 200, changed.data)
        student.refresh_from_db()
        self.assertEqual(student.bound_secretary_id, old_secretary)
        # The drawer submits a full form, including the unchanged secretary.
        saved = self.client.post('/api/schedule/adjust-group/', {'group_id': target.pk,
            'group_data': {'secretary': different_secretary.name, 'secretaryId': target_secretary,
                'date': '2026-10-12', 'timeRange': '14:00-17:00'}, 'expected_revision': 2}, format='json')
        self.assertEqual(saved.status_code, 200, saved.data)
        student.refresh_from_db()
        self.assertEqual(student.bound_secretary_id, old_secretary)

    def test_formal_time_form_does_not_rebind_a_moved_student(self):
        initial = self.linked()
        self.assertEqual(initial.status_code, 200, initial.data)
        formal = ScheduleVersion.objects.get(pk=initial.data['formal']['versionId'])
        source, target = list(formal.groups.select_related('secretary'))
        student = source.students.first()
        moved = self.client.post('/api/schedule/adjust/', {'action': 'move_student', 'student_id': student.pk,
            'from_group_id': source.pk, 'to_group_id': target.pk, 'expected_revision': 0}, format='json')
        self.assertEqual(moved.status_code, 200, moved.data)
        saved = self.client.post('/api/schedule/adjust-group/', {'group_id': target.pk,
            'group_data': {'secretary': target.secretary.name, 'secretaryId': target.secretary_id,
                'date': '2026-10-19', 'timeRange': '09:00-12:00'}, 'expected_revision': 1}, format='json')
        self.assertEqual(saved.status_code, 200, saved.data)
        formal.refresh_from_db()
        binding = next(s['secretary_id'] for s in formal.input_snapshot['students'] if s['id'] == student.pk)
        self.assertEqual(binding, source.secretary_id)
        current = self.client.get('/api/schedule/current/', {'defense_type': 'formal'})
        self.assertTrue(any(c['type'] == '秘书连续性提示' and c['level'] == 'error'
            for c in current.data['conflicts']), current.data['conflicts'])

    def test_explicit_formal_secretary_change_updates_only_its_version_binding(self):
        initial = self.linked()
        self.assertEqual(initial.status_code, 200, initial.data)
        formal = ScheduleVersion.objects.get(pk=initial.data['formal']['versionId'])
        group = formal.groups.first()
        new_secretary = self.teachers[-1]
        changed = self.client.post('/api/schedule/adjust/', {'action': 'change_secretary',
            'group_id': group.pk, 'new_secretary_id': new_secretary.pk, 'expected_revision': 0}, format='json')
        self.assertEqual(changed.status_code, 200, changed.data)
        formal.refresh_from_db()
        members = set(group.students.values_list('pk', flat=True))
        self.assertTrue(all(s['secretary_id'] == new_secretary.pk for s in formal.input_snapshot['students'] if s['id'] in members))
        self.assertFalse(any(c['type'] == 'secretary_continuity_broken' and group.pk in c.get('group_db_ids', [])
            for c in changed.data['conflicts']))

    def test_linked_pre_revalidation_keeps_new_round_secretary_bindings(self):
        initial = self.linked()
        self.assertEqual(initial.status_code, 200, initial.data)
        pre = ScheduleVersion.objects.get(pk=initial.data['pre']['versionId'])
        source, target = list(pre.groups.all())
        self.assertNotEqual(source.secretary_id, target.secretary_id)
        student = source.students.first()
        moved = self.client.post('/api/schedule/adjust/', {'action': 'move_student', 'student_id': student.pk,
            'from_group_id': source.pk, 'to_group_id': target.pk, 'expected_revision': 0}, format='json')
        self.assertEqual(moved.status_code, 200, moved.data)
        current = self.client.get('/api/schedule/current/', {'defense_type': 'pre'})
        self.assertTrue(any(c['type'] == '秘书连续性提示' and c['level'] == 'error'
            for c in current.data['conflicts']), current.data['conflicts'])

    def test_formal_validation_ignores_later_global_secretary_binding(self):
        initial = self.linked()
        self.assertEqual(initial.status_code, 200, initial.data)
        retired = Teacher.objects.create(name='后续停用的秘书', title='讲师', roles=['秘书'], is_active=False)
        Student.objects.update(bound_secretary=retired, secretary_name=retired.name)
        checked = self.client.post('/api/schedule/check-conflicts/', {'defense_type': 'formal'}, format='json')
        self.assertEqual(checked.status_code, 200, checked.data)
        self.assertFalse([c for c in checked.data if c['level'] == 'error'], checked.data)

    def test_formal_student_rebind_does_not_change_source_pre_binding(self):
        initial = self.linked()
        self.assertEqual(initial.status_code, 200, initial.data)
        formal = ScheduleVersion.objects.get(pk=initial.data['formal']['versionId'])
        source, target = list(formal.groups.all())
        student = source.students.first()
        original_secretary = student.bound_secretary_id
        moved = self.client.post('/api/schedule/adjust/', {'action': 'move_student', 'student_id': student.pk,
            'from_group_id': source.pk, 'to_group_id': target.pk, 'expected_revision': 0,
            'preserve_secretary': False, 'secretary_id': target.secretary_id}, format='json')
        self.assertEqual(moved.status_code, 200, moved.data)
        student.refresh_from_db()
        self.assertEqual(student.bound_secretary_id, original_secretary)
        formal.refresh_from_db()
        binding = next(s['secretary_id'] for s in formal.input_snapshot['students'] if s['id'] == student.pk)
        self.assertEqual(binding, target.secretary_id)
        self.assertEqual(moved.data['movement']['secretaryId'], target.secretary_id)

    def test_move_explicitly_uses_selected_secretary_without_rebinding_target_students(self):
        generated = self.client.post('/api/schedule/generate/', {'rules': self.rules('pre')}, format='json')
        self.assertEqual(generated.status_code, 200, generated.data)
        version = ScheduleVersion.objects.get(pk=generated.data['versionId'])
        source, target = list(version.groups.all())
        student = source.students.first()
        previous = dict(target.students.values_list('pk', 'bound_secretary_id'))
        replacement = self.teachers[-1]
        moved = self.client.post('/api/schedule/adjust/', {'action': 'move_student', 'student_id': student.pk,
            'from_group_id': source.pk, 'to_group_id': target.pk, 'expected_revision': 0,
            'preserve_secretary': False, 'secretary_id': replacement.pk}, format='json')
        self.assertEqual(moved.status_code, 200, moved.data)
        student.refresh_from_db()
        self.assertEqual(student.bound_secretary_id, replacement.pk)
        self.assertEqual(dict(target.students.exclude(pk=student.pk).values_list('pk', 'bound_secretary_id')), previous)

    def test_fk_snapshots_are_serializable_and_excel_contains_roster_remarks(self):
        generated = self.client.post('/api/schedule/generate/', {'rules': self.rules('pre')}, format='json')
        self.assertEqual(generated.status_code, 200, generated.data)
        version = ScheduleVersion.objects.get(pk=generated.data['versionId'])
        json.dumps(capture_export_groups(version))
        self.students[0].remark = '=SUM(1,2)'
        self.students[0].save()
        exported = self.client.get('/api/schedule/export/', {'defense_type': 'pre'})
        self.assertEqual(exported.status_code, 200)
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(exported.content))
        self.assertIn('教室学生明细', workbook.sheetnames)
        values = [cell for row in workbook['教室学生明细'] for cell in row if cell.value == '=SUM(1,2)']
        self.assertEqual(len(values), 1)
        self.assertEqual(values[0].data_type, 's')

    def test_publish_rechecks_v2_role_qualifications_after_reference_edit(self):
        generated = self.client.post('/api/schedule/generate/', {'rules': self.rules('pre')}, format='json')
        self.assertEqual(generated.status_code, 200, generated.data)
        version = ScheduleVersion.objects.get(pk=generated.data['versionId'])
        chair = version.groups.first().chair
        chair.title = '助教'
        chair.save()
        response = self.client.post('/api/schedule/publish/', {'version_id': version.pk, 'expected_revision': 0}, format='json')
        self.assertEqual(response.status_code, 400, response.data)
        self.assertTrue(any(c['type'] == '角色资格冲突' and c['level'] == 'error' for c in response.data['conflicts']))

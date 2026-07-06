import ast
import unittest
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, transaction
from django.test import TestCase
from rest_framework.test import APIClient

import algorithm
from .models import Group, Room, ScheduleVersion, Student, Teacher


def make_xlsx_upload(filename, rows):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(
        filename,
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


class BaseDataConstraintTests(TestCase):
    def test_teacher_name_is_unique_at_database_level(self):
        Teacher.objects.create(name='同名教师', college='计算机学院', title='教授')

        with self.assertRaises(IntegrityError), transaction.atomic():
            Teacher.objects.create(name='同名教师', college='软件学院', title='副教授')

    def test_student_no_is_unique_at_database_level(self):
        Student.objects.create(name='同名学生', student_no='3122158001', student_type='学硕', campus='创新港')

        with self.assertRaises(IntegrityError), transaction.atomic():
            Student.objects.create(name='另一学生', student_no='3122158001', student_type='专硕', campus='兴庆')

    def test_student_name_allows_duplicates_with_different_student_no(self):
        Student.objects.create(name='同名学生', student_no='3122158001', student_type='学硕', campus='创新港')

        Student.objects.create(name='同名学生', student_no='3122158002', student_type='专硕', campus='兴庆')

        self.assertEqual(Student.objects.filter(name='同名学生').count(), 2)

    def test_room_name_is_unique_per_campus_at_database_level(self):
        Room.objects.create(campus='创新港', name='A101', capacity=30)

        with self.assertRaises(IntegrityError), transaction.atomic():
            Room.objects.create(campus='创新港', name='A101', capacity=40)

    def test_room_name_can_repeat_across_campuses_at_database_level(self):
        Room.objects.create(campus='创新港', name='A101', capacity=30)

        Room.objects.create(campus='兴庆', name='A101', capacity=40)

        self.assertEqual(Room.objects.filter(name='A101').count(), 2)


class DesktopStartupDependencyTests(TestCase):
    def test_api_views_defers_heavy_spreadsheet_imports_until_needed(self):
        source = Path(__file__).with_name('views.py').read_text(encoding='utf-8')
        tree = ast.parse(source)
        top_level_imports = set()
        for node in tree.body:
            if isinstance(node, ast.Import):
                top_level_imports.update(alias.name.split('.')[0] for alias in node.names)
            elif isinstance(node, ast.ImportFrom) and node.module:
                top_level_imports.add(node.module.split('.')[0])

        self.assertNotIn('pandas', top_level_imports)
        self.assertNotIn('openpyxl', top_level_imports)


class AuthContractTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user(username='admin', password='strong-pass-123')
        Teacher.objects.create(name='教师1', college='计算机学院', title='教授')

    def test_api_requires_authentication_except_login(self):
        response = self.client.get('/api/teachers/')

        self.assertEqual(response.status_code, 401)

        login_response = self.client.post(
            '/api/auth/login/',
            {'username': 'admin', 'password': 'strong-pass-123'},
            format='json',
        )
        self.assertEqual(login_response.status_code, 200)
        self.assertIn('token', login_response.data)
        self.assertFalse(login_response.data['isAdmin'])

        self.client.credentials(HTTP_AUTHORIZATION=f"Token {login_response.data['token']}")
        authed_response = self.client.get('/api/teachers/')
        self.assertEqual(authed_response.status_code, 200)
        self.assertEqual(len(authed_response.data), 1)


class PermissionContractTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user(username='viewer', password='strong-pass-123')
        self.client.force_authenticate(self.user)
        Teacher.objects.create(name='教师1', college='计算机学院', title='教授')

    def assert_forbidden(self, response):
        self.assertEqual(response.status_code, 403)

    def test_regular_user_can_read_reference_data_and_schedule_state(self):
        self.assertEqual(self.client.get('/api/teachers/').status_code, 200)
        self.assertEqual(self.client.get('/api/students/').status_code, 200)
        self.assertEqual(self.client.get('/api/rooms/').status_code, 200)
        self.assertEqual(self.client.get('/api/rule-config/', {'defense_type': 'pre'}).status_code, 200)
        self.assertEqual(self.client.get('/api/operation-logs/').status_code, 200)
        self.assertEqual(self.client.get('/api/schedule/current/', {'defense_type': 'pre'}).status_code, 200)
        self.assertEqual(
            self.client.post('/api/schedule/check-conflicts/', {'defense_type': 'pre'}, format='json').status_code,
            200,
        )

    def test_regular_user_cannot_modify_operational_data(self):
        unsafe_requests = [
            self.client.post(
                '/api/teachers/',
                {'name': '教师2', 'college': '计算机学院', 'title': '副教授'},
                format='json',
            ),
            self.client.post('/api/teachers/import_data/', {}, format='multipart'),
            self.client.post('/api/teachers/batch_delete/', {'ids': [1]}, format='json'),
            self.client.post('/api/rule-config/', {'defense_type': 'pre', 'defenseType': '预答辩'}, format='json'),
            self.client.post('/api/schedule/generate/', {'rules': {'defense_type': 'pre'}}, format='json'),
            self.client.post('/api/schedule/adjust-group/', {'group_id': 1, 'group_data': {'groupName': 'G1'}}, format='json'),
            self.client.delete('/api/operation-logs/'),
        ]

        for response in unsafe_requests:
            self.assert_forbidden(response)

    def test_admin_user_can_modify_operational_data(self):
        admin = User.objects.create_superuser(username='admin', password='strong-pass-123')
        self.client.force_authenticate(admin)

        response = self.client.post(
            '/api/teachers/',
            {'name': '教师2', 'college': '计算机学院', 'title': '副教授'},
            format='json',
        )

        self.assertEqual(response.status_code, 201)


class IntegrationContractTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_superuser(username='admin', password='strong-pass-123')
        self.client.force_authenticate(self.user)

    def test_rule_config_api_supports_frontend_contract(self):
        response = self.client.get('/api/rule-config/', {'defense_type': 'pre'})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['defenseType'], '预答辩')
        self.assertIn('studentCount', response.data)
        self.assertEqual(response.data['endDate'], '2025-05-20')

        response = self.client.post(
            '/api/rule-config/',
            {
                'defense_type': 'pre',
                'defenseType': '预答辩',
                'startDate': '2025-05-10',
                'endDate': '2025-05-12',
                'studentCount': {'target': 6, 'min': 3, 'max': 8},
            },
            format='json',
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['studentCount']['target'], 6)
        self.assertEqual(response.data['endDate'], '2025-05-12')

    def test_rule_config_rejects_end_date_before_start_date(self):
        response = self.client.post(
            '/api/rule-config/',
            {
                'defense_type': 'pre',
                'defenseType': '预答辩',
                'startDate': '2025-05-12',
                'endDate': '2025-05-10',
            },
            format='json',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('endDate', response.data)

    def test_rule_config_rejects_invalid_student_count_bounds(self):
        response = self.client.post(
            '/api/rule-config/',
            {
                'defense_type': 'pre',
                'defenseType': '预答辩',
                'studentCount': {'target': 6, 'min': 7, 'max': 8},
            },
            format='json',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('studentCount', response.data)

    def test_operation_logs_api_supports_frontend_contract(self):
        create_response = self.client.post(
            '/api/operation-logs/',
            {
                'type': 'save_rule_config',
                'module': 'rule-config',
                'description': 'updated rule config',
                'operator': 'admin',
                'result': 'success',
            },
            format='json',
        )

        self.assertEqual(create_response.status_code, 201)
        self.assertIn('createdAt', create_response.data)

        list_response = self.client.get('/api/operation-logs/')
        self.assertEqual(list_response.status_code, 200)
        self.assertEqual(len(list_response.data), 1)

        clear_response = self.client.delete('/api/operation-logs/')
        self.assertEqual(clear_response.status_code, 200)
        self.assertEqual(clear_response.data['message'], '日志已清空')
        self.assertEqual(self.client.get('/api/operation-logs/').data, [])

    def test_import_rejects_oversized_files_before_parsing(self):
        oversized_file = SimpleUploadedFile(
            'teachers.csv',
            b'name,title\n' + (b'x' * (5 * 1024 * 1024 + 1)),
            content_type='text/csv',
        )

        response = self.client.post('/api/teachers/import_data/', {'file': oversized_file})

        self.assertEqual(response.status_code, 400)
        self.assertIn('文件大小不能超过', response.data['error'])
        self.assertEqual(Teacher.objects.count(), 0)

    def test_import_is_atomic_and_returns_structured_row_errors(self):
        upload = SimpleUploadedFile(
            'teachers.csv',
            'name,college,title\n有效教师,计算机学院,教授\n无职称教师,计算机学院,\n'.encode('utf-8'),
            content_type='text/csv',
        )

        response = self.client.post('/api/teachers/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 400)
        self.assertEqual(Teacher.objects.count(), 0)
        self.assertEqual(response.data['errors'][0]['row'], 3)
        self.assertIn('title', response.data['errors'][0]['errors'])

    def test_import_rejects_empty_data_files(self):
        upload = SimpleUploadedFile(
            'teachers.csv',
            'name,title\n'.encode('utf-8'),
            content_type='text/csv',
        )

        response = self.client.post('/api/teachers/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 400)
        self.assertIn('没有可导入的数据', response.data['error'])
        self.assertEqual(Teacher.objects.count(), 0)

    def test_import_rejects_files_with_too_many_rows(self):
        csv_text = 'name,title\n' + ''.join(f'教师{i},教授\n' for i in range(1001))
        upload = SimpleUploadedFile(
            'teachers.csv',
            csv_text.encode('utf-8'),
            content_type='text/csv',
        )

        response = self.client.post('/api/teachers/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 400)
        self.assertIn('单次最多导入', response.data['error'])
        self.assertEqual(Teacher.objects.count(), 0)

    def test_import_accepts_supported_extensions_case_insensitively(self):
        upload = SimpleUploadedFile(
            'TEACHERS.CSV',
            'name,title\n教师1,教授\n'.encode('utf-8'),
            content_type='text/csv',
        )

        response = self.client.post('/api/teachers/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(Teacher.objects.count(), 1)

    def test_teacher_api_rejects_duplicate_names(self):
        Teacher.objects.create(name='重复教师', college='计算机学院', title='教授')

        response = self.client.post(
            '/api/teachers/',
            {'name': '重复教师', 'college': '软件学院', 'title': '副教授'},
            format='json',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('name', response.data)
        self.assertEqual(Teacher.objects.filter(name='重复教师').count(), 1)

    def test_import_rejects_duplicate_teacher_names_in_same_file(self):
        upload = SimpleUploadedFile(
            'teachers.csv',
            'name,title\n重复教师,教授\n重复教师,副教授\n'.encode('utf-8'),
            content_type='text/csv',
        )

        response = self.client.post('/api/teachers/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 400)
        self.assertEqual(Teacher.objects.count(), 0)
        self.assertEqual(response.data['errors'][0]['row'], 3)
        self.assertIn('name', response.data['errors'][0]['errors'])

    def test_student_api_rejects_duplicate_names(self):
        Student.objects.create(name='重复学生', student_type='学硕', campus='创新港')

        response = self.client.post(
            '/api/students/',
            {'name': '重复学生', 'studentType': '专硕', 'campus': '兴庆'},
            format='json',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('name', response.data)
        self.assertEqual(Student.objects.filter(name='重复学生').count(), 1)

    def test_room_api_rejects_duplicate_names_in_same_campus(self):
        Room.objects.create(campus='创新港', name='A101', capacity=30)

        response = self.client.post(
            '/api/rooms/',
            {'campus': '创新港', 'name': 'A101', 'capacity': 40},
            format='json',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('name', response.data)
        self.assertEqual(Room.objects.filter(campus='创新港', name='A101').count(), 1)

    def test_room_api_allows_same_name_in_different_campus(self):
        Room.objects.create(campus='创新港', name='A101', capacity=30)

        response = self.client.post(
            '/api/rooms/',
            {'campus': '兴庆', 'name': 'A101', 'capacity': 40},
            format='json',
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(Room.objects.filter(name='A101').count(), 2)

    def test_import_rejects_duplicate_student_names_in_same_file(self):
        upload = SimpleUploadedFile(
            'students.csv',
            'name,studentType,campus\n重复学生,学硕,创新港\n重复学生,专硕,兴庆\n'.encode('utf-8'),
            content_type='text/csv',
        )

        response = self.client.post('/api/students/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 400)
        self.assertEqual(Student.objects.count(), 0)
        self.assertEqual(response.data['errors'][0]['row'], 3)
        self.assertIn('name', response.data['errors'][0]['errors'])

    def test_import_rejects_duplicate_rooms_in_same_campus_in_same_file(self):
        upload = SimpleUploadedFile(
            'rooms.csv',
            'campus,name,capacity\n创新港,A101,30\n创新港,A101,40\n'.encode('utf-8'),
            content_type='text/csv',
        )

        response = self.client.post('/api/rooms/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 400)
        self.assertEqual(Room.objects.count(), 0)
        self.assertEqual(response.data['errors'][0]['row'], 3)
        self.assertIn('name', response.data['errors'][0]['errors'])

    def test_teacher_import_accepts_chinese_headers_and_separators(self):
        upload = SimpleUploadedFile(
            'teachers.csv',
            (
                '教师姓名,所属学院,职称,是否外院,可担任角色,参加答辩类型\n'
                '张老师,计算机学院,教授,否,主席，普通专家,预答辩；正式答辩\n'
            ).encode('utf-8'),
            content_type='text/csv',
        )

        response = self.client.post('/api/teachers/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 200)
        teacher = Teacher.objects.get(name='张老师')
        self.assertEqual(teacher.college, '计算机学院')
        self.assertEqual(teacher.roles, ['主席', '普通专家'])
        self.assertEqual(teacher.available_types, ['预答辩', '正式答辩'])

    def test_student_import_accepts_chinese_headers(self):
        upload = SimpleUploadedFile(
            'students.csv',
            (
                '学生姓名,学生类型,导师姓名,所属校区,参加答辩类型,对应秘书姓名\n'
                '学生甲,学硕,张老师,创新港,预答辩，正式答辩,秘书王\n'
            ).encode('utf-8'),
            content_type='text/csv',
        )

        response = self.client.post('/api/students/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 200)
        student = Student.objects.get(name='学生甲')
        self.assertEqual(student.student_type, '学硕')
        self.assertEqual(student.mentor_name, '张老师')
        self.assertEqual(student.defense_types, ['预答辩', '正式答辩'])
        self.assertEqual(student.secretary_name, '秘书王')

    def test_room_import_accepts_chinese_headers(self):
        upload = SimpleUploadedFile(
            'rooms.csv',
            '校区,教室名称,容量,可用时间段\n创新港,A101,30,2025-05-10 09:00-12:00\n'.encode('utf-8'),
            content_type='text/csv',
        )

        response = self.client.post('/api/rooms/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 200)
        room = Room.objects.get(name='A101')
        self.assertEqual(room.campus, '创新港')
        self.assertEqual(room.capacity, 30)
        self.assertEqual(room.available_times, '2025-05-10 09:00-12:00')

    def test_teacher_import_accepts_assigned_leader_secretary_roster(self):
        upload = make_xlsx_upload(
            '指定组长秘书.xlsx',
            [
                ['序号', '姓名', '所在学院', '性别', '职称', None],
                [1, '宋永红', '软件学院', '女', '研究员', '组长/主席'],
                [2, '田暄', '软件学院', '男', '工程师', '秘书'],
            ],
        )

        response = self.client.post('/api/teachers/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 200)
        leader = Teacher.objects.get(name='宋永红')
        secretary = Teacher.objects.get(name='田暄')
        self.assertEqual(leader.college, '软件学院')
        self.assertEqual(leader.roles, ['组长', '主席'])
        self.assertEqual(secretary.roles, ['秘书'])

    def test_student_import_accepts_midterm_location_sheet_with_header_after_blank_row(self):
        upload = make_xlsx_upload(
            '22级中期考核答辩地点.xlsx',
            [
                [None, None, None, None, None, None, None, None],
                ['序号', '学号', '姓名', '性别', '学科', '导师', '答辩地点', '答辩地点填：兴庆/创新港'],
                [1, '3122158001', '崔东森', '男', '计算机科学与技术', '王晨旭', '创新港', None],
            ],
        )

        response = self.client.post('/api/students/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 200)
        student = Student.objects.get(name='崔东森')
        self.assertEqual(student.student_no, '3122158001')
        self.assertEqual(student.gender, '男')
        self.assertEqual(student.student_type, '计算机科学与技术')
        self.assertEqual(student.mentor_name, '王晨旭')
        self.assertEqual(student.campus, '创新港')
        self.assertEqual(student.defense_types, ['中期答辩'])

    def test_student_import_accepts_formal_location_sheet_with_title_rows(self):
        upload = make_xlsx_upload(
            '23级答辩地点统计-含学生信息-导师信息-答辩地点.xlsx',
            [
                ['软件工程硕士23级答辩地点统计', None, None, None, None, None, None, None, None],
                ['2023年9月入学', None, None, None, None, None, None, None, None],
                ['序号', '学号', '姓名', '性别', '学科', '导师', '答辩地点', '备注', '答辩地点填：兴庆/创新港'],
                [1, '3123158001', '卓佳麟', '男', '计算机科学与技术', '王志', '创新港', None, None],
            ],
        )

        response = self.client.post('/api/students/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 200)
        student = Student.objects.get(name='卓佳麟')
        self.assertEqual(student.student_no, '3123158001')
        self.assertEqual(student.gender, '男')
        self.assertEqual(student.student_type, '计算机科学与技术')
        self.assertEqual(student.mentor_name, '王志')
        self.assertEqual(student.campus, '创新港')
        self.assertEqual(student.defense_types, ['正式答辩'])

    def test_student_import_normalizes_numeric_student_no_cells(self):
        # Excel 数值单元格的学号会被读成 int，必须归一化为不带小数尾巴的字符串
        upload = make_xlsx_upload(
            '23级答辩地点统计.xlsx',
            [
                ['序号', '学号', '姓名', '性别', '学科', '导师', '答辩地点'],
                [1, 3123158777, '数值学号学生', '女', '软件工程', '王志', '创新港'],
            ],
        )

        response = self.client.post('/api/students/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 200)
        student = Student.objects.get(name='数值学号学生')
        self.assertEqual(student.student_no, '3123158777')

    def test_student_import_allows_same_name_students_with_different_student_no(self):
        upload = make_xlsx_upload(
            '23级答辩地点统计.xlsx',
            [
                ['序号', '学号', '姓名', '性别', '学科', '导师', '答辩地点'],
                [1, '3123158801', '重名学生', '男', '软件工程', '王志', '创新港'],
                [2, '3123158802', '重名学生', '女', '软件工程', '李明', '兴庆'],
            ],
        )

        response = self.client.post('/api/students/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 200)
        students = Student.objects.filter(name='重名学生').order_by('student_no')
        self.assertEqual(students.count(), 2)
        self.assertEqual([s.student_no for s in students], ['3123158801', '3123158802'])
        self.assertEqual([s.mentor_name for s in students], ['王志', '李明'])

    def test_student_import_merges_defense_types_for_existing_real_data_rows(self):
        Student.objects.create(
            name='同名学生',
            student_type='计算机科学与技术',
            mentor_name='旧导师',
            campus='创新港',
            defense_types=['中期答辩'],
        )
        upload = make_xlsx_upload(
            '23级答辩地点统计-含学生信息-导师信息-答辩地点.xlsx',
            [
                ['序号', '学号', '姓名', '性别', '学科', '导师', '答辩地点'],
                [1, '3123158999', '同名学生', '男', '计算机科学与技术', '新导师', '兴庆'],
            ],
        )

        response = self.client.post('/api/students/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 200)
        student = Student.objects.get(name='同名学生')
        self.assertEqual(student.mentor_name, '新导师')
        self.assertEqual(student.campus, '兴庆')
        self.assertEqual(student.defense_types, ['中期答辩', '正式答辩'])

    def test_student_import_merges_explicit_defense_type_text_for_existing_rows(self):
        Student.objects.create(
            name='已有学生',
            student_type='学硕',
            mentor_name='旧导师',
            campus='创新港',
            defense_types=['中期答辩'],
        )
        upload = SimpleUploadedFile(
            'students.csv',
            (
                '学生姓名,学生类型,导师姓名,所属校区,参加答辩类型\n'
                '已有学生,专硕,新导师,兴庆,预答辩，正式答辩\n'
            ).encode('utf-8'),
            content_type='text/csv',
        )

        response = self.client.post('/api/students/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 200)
        student = Student.objects.get(name='已有学生')
        self.assertEqual(student.student_type, '专硕')
        self.assertEqual(student.mentor_name, '新导师')
        self.assertEqual(student.campus, '兴庆')
        self.assertEqual(student.defense_types, ['中期答辩', '预答辩', '正式答辩'])

    def test_room_import_accepts_borrow_request_sheet_with_date_and_class_periods(self):
        upload = make_xlsx_upload(
            '教室借用申请-兴庆教室.xlsx',
            [
                ['审核状态', '借用人姓名', '校区', '教学楼', '教室名称', '使用日期', '使用时间'],
                ['已通过', '冯硕', '兴庆校区', '中2楼', '中2-1209', '2024-03-28', '第1节-第4节'],
            ],
        )

        response = self.client.post('/api/rooms/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 200)
        room = Room.objects.get(name='中2-1209')
        self.assertEqual(room.campus, '兴庆')
        self.assertEqual(room.capacity, 30)
        self.assertEqual(room.available_times, '2024-03-28 08:00-12:00')

    def test_room_import_accepts_card_style_innovation_harbor_borrow_sheet(self):
        upload = make_xlsx_upload(
            '教室借用申请-创新港教室.xlsx',
            [
                ['预答辩(不公开)', None, None, None],
                ['状态：正在审核单号:101990393详情', None, None, None],
                ['4-3214 电子与信息学部 等待审核 2026-04-10 09:00至13:00 服务： 撤销 | 评价', None, None, None],
                ['详细 | 复制申请', None, '审核通过', None],
            ],
        )

        response = self.client.post('/api/rooms/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 200)
        room = Room.objects.get(name='4-3214')
        self.assertEqual(room.campus, '创新港')
        self.assertEqual(room.capacity, 30)
        self.assertEqual(room.available_times, '2026-04-10 09:00-13:00')

    def test_room_import_skips_card_style_noise_rows_and_merges_duplicate_rooms(self):
        upload = make_xlsx_upload(
            '教室借用申请-创新港教室.xlsx',
            [
                ['预答辩(不公开)', None, None, None, None, None],
                ['状态：正在审核单号:101990393详情', None, None, None, None, None],
                ['4-3214', '电子与信息学部', '等待审核', '2026-04-10 09:00至13:00', '服务：', '撤销 | 评价'],
                ['详细 | 复制申请', None, '审核通过', None, None, None],
                ['预答辩(不公开)', None, None, None, None, None],
                ['状态：正在审核单号:101990363详情', None, None, None, None, None],
                ['4-3214', '电子与信息学部', '等待审核', '2026-04-09 08:50至18:10', '服务：', '撤销 | 评价'],
            ],
        )

        response = self.client.post('/api/rooms/import_data/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(Room.objects.count(), 1)
        room = Room.objects.get(name='4-3214')
        self.assertEqual(room.campus, '创新港')
        self.assertEqual(
            room.available_times,
            '2026-04-10 09:00-13:00,2026-04-09 08:50-18:10',
        )


class ScheduleContractTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_superuser(username='admin', password='strong-pass-123')
        self.client.force_authenticate(self.user)
        for index, title in enumerate(['教授', '副教授', '讲师', '讲师'], start=1):
            Teacher.objects.create(
                name=f'教师{index}',
                college='计算机学院',
                title=title,
                roles=['专家', '秘书'],
                available_types=['预答辩'],
            )
        Room.objects.create(
            campus='创新港',
            name='A101',
            capacity=30,
            available_times='',
        )
        for index in range(1, 4):
            Student.objects.create(
                name=f'学生{index}',
                student_type='学硕',
                mentor_name='教师1',
                campus='创新港',
                defense_types=['预答辩'],
                secretary_name='教师4',
            )

    def _create_manual_adjustment_group(self):
        teacher = Teacher.objects.first()
        secretary = Teacher.objects.last()
        room = Room.objects.first()
        student = Student.objects.first()
        version = ScheduleVersion.objects.create(
            version=1,
            defense_type='pre',
            is_current=True,
            rules_snapshot={},
        )
        group = Group.objects.create(
            schedule_version=version,
            group_id='G1',
            time='2025-05-10 09:00-12:00',
            room=room,
            campus='创新港',
            chair=teacher,
            secretary=secretary,
        )
        group.experts.add(teacher)
        group.students.add(student)
        return group

    def test_schedule_generate_current_conflict_check_and_export_use_current_models(self):
        generate_response = self.client.post(
            '/api/schedule/generate/',
            {
                'rules': {
                    'defense_type': 'pre',
                    'start_date': '2025-05-10',
                    'end_date': '2025-05-10',
                    'group_size': 3,
                    'expert_count': 2,
                    'avoid_weekend': False,
                    'avoid_supervisor': False,
                }
            },
            format='json',
        )

        self.assertEqual(generate_response.status_code, 200)
        self.assertEqual(generate_response.data['defenseType'], '预答辩')
        self.assertEqual(generate_response.data['groups'][0]['students'][0]['studentType'], '学硕')
        self.assertEqual(generate_response.data['groups'][0]['students'][0]['mentorName'], '教师1')

        current_response = self.client.get('/api/schedule/current/', {'defense_type': 'pre'})
        self.assertEqual(current_response.status_code, 200)
        self.assertEqual(len(current_response.data['groups']), 1)

        conflicts_response = self.client.post(
            '/api/schedule/check-conflicts/',
            {'defense_type': 'pre'},
            format='json',
        )
        self.assertEqual(conflicts_response.status_code, 200)
        self.assertIsInstance(conflicts_response.data, list)

        export_response = self.client.get('/api/schedule/export/', {'defense_type': 'pre'})
        self.assertEqual(export_response.status_code, 200)
        self.assertEqual(
            export_response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    def test_schedule_generation_maps_frontend_model_fields_for_constraints(self):
        Teacher.objects.all().delete()
        Student.objects.all().delete()
        Room.objects.all().delete()
        mentor = Teacher.objects.create(
            name='导师张',
            college='计算机学院',
            title='教授',
            roles=['专家'],
            available_types=['预答辩'],
        )
        eligible = Teacher.objects.create(
            name='专家李',
            college='计算机学院',
            title='副教授',
            roles=['专家'],
            available_types=['预答辩'],
        )
        Teacher.objects.create(
            name='秘书王',
            college='计算机学院',
            title='讲师',
            roles=['秘书'],
            available_types=['预答辩'],
        )
        Room.objects.create(campus='创新港', name='A101', capacity=30, available_times='')
        Student.objects.create(
            name='学生1',
            student_type='学硕',
            mentor_name=mentor.name,
            campus='创新港',
            defense_types=['预答辩'],
            secretary_name='秘书王',
        )

        response = self.client.post(
            '/api/schedule/generate/',
            {
                'rules': {
                    'defense_type': 'pre',
                    'start_date': '2025-05-10',
                    'end_date': '2025-05-10',
                    'group_size': 1,
                    'expert_count': 1,
                    'avoid_weekend': False,
                    'avoid_supervisor': True,
                }
            },
            format='json',
        )

        self.assertEqual(response.status_code, 200)
        group = response.data['groups'][0]
        self.assertNotIn(mentor.name, [teacher['name'] for teacher in group['teachers']])
        self.assertIn(eligible.name, [teacher['name'] for teacher in group['teachers']])

    def test_frontend_adjust_and_export_aliases_are_supported(self):
        teacher = Teacher.objects.first()
        secretary = Teacher.objects.last()
        room = Room.objects.first()
        student = Student.objects.first()
        version = ScheduleVersion.objects.create(
            version=1,
            defense_type='pre',
            is_current=True,
            rules_snapshot={},
        )
        group = Group.objects.create(
            schedule_version=version,
            group_id='G1',
            time='2025-05-10 09:00-12:00',
            room=room,
            campus='创新港',
            chair=teacher,
            secretary=secretary,
        )
        group.experts.add(teacher)
        group.students.add(student)

        adjust_response = self.client.post(
            '/api/schedule/adjust-group/',
            {
                'defense_type': 'pre',
                'group_id': group.id,
                'group_data': {
                    'groupName': 'G1-调整',
                    'date': '2025-05-11',
                    'timeRange': '14:00-16:00',
                    'campus': '兴庆',
                    'classroom': room.name,
                    'chairman': teacher.name,
                    'secretary': secretary.name,
                    'teachers': [{'id': teacher.id, 'name': teacher.name, 'title': teacher.title, 'roles': teacher.roles}],
                    'students': [{'id': student.id, 'name': student.name, 'studentType': student.student_type, 'mentorName': student.mentor_name}],
                },
            },
            format='json',
        )
        self.assertEqual(adjust_response.status_code, 200)
        group.refresh_from_db()
        self.assertEqual(group.group_id, 'G1-调整')
        self.assertEqual(group.time, '2025-05-11 14:00-16:00')

        export_response = self.client.get('/api/schedule/export-excel/', {'defenseType': '预答辩'})
        self.assertEqual(export_response.status_code, 200)
        self.assertEqual(
            export_response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    def test_adjust_group_rejects_unknown_references_without_partial_update(self):
        teacher = Teacher.objects.first()
        secretary = Teacher.objects.last()
        room = Room.objects.first()
        student = Student.objects.first()
        version = ScheduleVersion.objects.create(
            version=1,
            defense_type='pre',
            is_current=True,
            rules_snapshot={},
        )
        group = Group.objects.create(
            schedule_version=version,
            group_id='G1',
            time='2025-05-10 09:00-12:00',
            room=room,
            campus='创新港',
            chair=teacher,
            secretary=secretary,
        )
        group.experts.add(teacher)
        group.students.add(student)

        response = self.client.post(
            '/api/schedule/adjust-group/',
            {
                'group_id': group.id,
                'group_data': {
                    'groupName': '不应保存',
                    'classroom': '不存在的教室',
                    'chairman': teacher.name,
                    'secretary': secretary.name,
                    'teachers': [{'id': 999999, 'name': '不存在的教师'}],
                    'students': [{'id': student.id, 'name': student.name}],
                },
            },
            format='json',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('教室不存在', response.data['error'])
        group.refresh_from_db()
        self.assertEqual(group.group_id, 'G1')
        self.assertEqual(list(group.experts.values_list('id', flat=True)), [teacher.id])

    def test_adjust_group_rejects_invalid_time_without_partial_update(self):
        group = self._create_manual_adjustment_group()
        teacher = group.chair
        secretary = group.secretary
        room = group.room
        student = group.students.first()

        response = self.client.post(
            '/api/schedule/adjust-group/',
            {
                'group_id': group.id,
                'group_data': {
                    'groupName': '不应保存',
                    'date': '2025-05-11',
                    'timeRange': '16:00-14:00',
                    'classroom': room.name,
                    'chairman': teacher.name,
                    'secretary': secretary.name,
                    'teachers': [{'id': teacher.id, 'name': teacher.name}],
                    'students': [{'id': student.id, 'name': student.name}],
                },
            },
            format='json',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('时间格式无效', response.data['error'])
        group.refresh_from_db()
        self.assertEqual(group.group_id, 'G1')
        self.assertEqual(group.time, '2025-05-10 09:00-12:00')

    def test_adjust_action_rejects_invalid_time_without_partial_update(self):
        group = self._create_manual_adjustment_group()

        response = self.client.post(
            '/api/schedule/adjust/',
            {
                'action': 'change_time',
                'group_id': group.id,
                'new_time': 'not-a-time',
            },
            format='json',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('时间格式无效', response.data['error'])
        group.refresh_from_db()
        self.assertEqual(group.time, '2025-05-10 09:00-12:00')

    def test_adjust_action_rejects_missing_room_without_partial_update(self):
        group = self._create_manual_adjustment_group()
        original_room_id = group.room_id

        response = self.client.post(
            '/api/schedule/adjust/',
            {
                'action': 'change_room',
                'group_id': group.id,
                'new_room_id': 999999,
            },
            format='json',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('教室不存在', response.data['error'])
        group.refresh_from_db()
        self.assertEqual(group.room_id, original_room_id)

    def test_adjust_action_rejects_missing_chair_without_partial_update(self):
        group = self._create_manual_adjustment_group()
        original_chair_id = group.chair_id

        response = self.client.post(
            '/api/schedule/adjust/',
            {
                'action': 'change_chair',
                'group_id': group.id,
                'new_chair_id': 999999,
            },
            format='json',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('主席不存在', response.data['error'])
        group.refresh_from_db()
        self.assertEqual(group.chair_id, original_chair_id)

    def test_adjust_action_rejects_missing_secretary_without_partial_update(self):
        group = self._create_manual_adjustment_group()
        original_secretary_id = group.secretary_id

        response = self.client.post(
            '/api/schedule/adjust/',
            {
                'action': 'change_secretary',
                'group_id': group.id,
                'new_secretary_id': 999999,
            },
            format='json',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('秘书不存在', response.data['error'])
        group.refresh_from_db()
        self.assertEqual(group.secretary_id, original_secretary_id)

    def test_adjust_action_rejects_missing_expert_without_partial_update(self):
        group = self._create_manual_adjustment_group()
        original_expert_ids = list(group.experts.values_list('id', flat=True))

        response = self.client.post(
            '/api/schedule/adjust/',
            {
                'action': 'change_expert',
                'group_id': group.id,
                'old_expert_id': original_expert_ids[0],
                'new_expert_id': 999999,
            },
            format='json',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('教师不存在', response.data['error'])
        group.refresh_from_db()
        self.assertEqual(list(group.experts.values_list('id', flat=True)), original_expert_ids)

    def test_adjust_action_rejects_missing_student_without_partial_update(self):
        group = self._create_manual_adjustment_group()
        target_group = self._create_manual_adjustment_group()
        target_group.group_id = 'G2'
        target_group.save()
        original_student_ids = list(group.students.values_list('id', flat=True))
        target_student_ids = list(target_group.students.values_list('id', flat=True))

        response = self.client.post(
            '/api/schedule/adjust/',
            {
                'action': 'move_student',
                'student_id': 999999,
                'from_group_id': group.id,
                'to_group_id': target_group.id,
            },
            format='json',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('学生不存在', response.data['error'])
        group.refresh_from_db()
        target_group.refresh_from_db()
        self.assertEqual(list(group.students.values_list('id', flat=True)), original_student_ids)
        self.assertEqual(list(target_group.students.values_list('id', flat=True)), target_student_ids)


class ScheduleConflictContractTests(TestCase):
    """生成排期时的冲突链路与规则键契约"""

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_superuser(username='admin', password='strong-pass-123')
        self.client.force_authenticate(self.user)

    def _generate(self, rules):
        return self.client.post('/api/schedule/generate/', {'rules': rules}, format='json')

    def test_generate_persists_and_returns_frontend_shaped_conflicts(self):
        # 唯一的老师同时是学生导师，且开启导师回避 → 必然缺专家、缺秘书
        mentor = Teacher.objects.create(name='导师张', college='计算机学院', title='教授')
        Room.objects.create(campus='创新港', name='A101', capacity=30, available_times='')
        Student.objects.create(
            name='学生1', student_type='学硕', mentor_name=mentor.name,
            campus='创新港', defense_types=['预答辩'],
        )

        response = self._generate({
            'defense_type': 'pre',
            'start_date': '2025-05-10',
            'end_date': '2025-05-10',
            'group_size': 1,
            'expert_count': 1,
            'avoid_weekend': False,
            'avoid_supervisor': True,
        })

        self.assertEqual(response.status_code, 200)
        conflicts = response.data['conflicts']
        self.assertTrue(conflicts, '数据不足时生成结果必须携带冲突提示')
        for conflict in conflicts:
            self.assertIn('type', conflict)
            self.assertIn(conflict['level'], ['error', 'warning', 'info'])
            self.assertTrue(conflict['reason'])
            self.assertEqual(conflict['defenseType'], '预答辩')
        self.assertIn('人员冲突', [conflict['type'] for conflict in conflicts])

        version = ScheduleVersion.objects.get(defense_type='pre', is_current=True)
        self.assertTrue(version.conflicts_snapshot, '冲突快照必须持久化')

        current_response = self.client.get('/api/schedule/current/', {'defense_type': 'pre'})
        self.assertEqual(current_response.status_code, 200)
        self.assertEqual(current_response.data['conflicts'], conflicts)

    def test_generate_normalizes_mentor_avoidance_alias(self):
        mentor = Teacher.objects.create(name='导师张', college='计算机学院', title='教授')
        Teacher.objects.create(name='专家李', college='计算机学院', title='副教授')
        Teacher.objects.create(name='秘书王', college='计算机学院', title='讲师')
        Room.objects.create(campus='创新港', name='A101', capacity=30, available_times='')
        Student.objects.create(
            name='学生1', student_type='学硕', mentor_name=mentor.name,
            campus='创新港', defense_types=['预答辩'],
        )

        # 旧键名 mentor_avoidance 也必须触发导师回避
        response = self._generate({
            'defense_type': 'pre',
            'start_date': '2025-05-10',
            'end_date': '2025-05-10',
            'group_size': 1,
            'expert_count': 1,
            'avoid_weekend': False,
            'mentor_avoidance': True,
        })

        self.assertEqual(response.status_code, 200)
        version = ScheduleVersion.objects.get(defense_type='pre', is_current=True)
        self.assertTrue(version.rules_snapshot.get('avoid_supervisor'))
        group = response.data['groups'][0]
        self.assertNotIn(mentor.name, [teacher['name'] for teacher in group['teachers']])

    def test_generate_assigns_chair_when_need_chair_rule_sent(self):
        Teacher.objects.create(name='教授A', college='计算机学院', title='教授')
        Teacher.objects.create(name='讲师B', college='计算机学院', title='讲师')
        Teacher.objects.create(name='讲师C', college='计算机学院', title='讲师')
        Room.objects.create(campus='创新港', name='A101', capacity=30, available_times='')
        Student.objects.create(
            name='学生1', student_type='学硕', mentor_name='',
            campus='创新港', defense_types=['预答辩'],
        )

        response = self._generate({
            'defense_type': 'pre',
            'start_date': '2025-05-10',
            'end_date': '2025-05-10',
            'group_size': 1,
            'expert_count': 1,
            'avoid_weekend': False,
            'avoid_supervisor': False,
            'need_chair': True,
            'chair_title': '教授',
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['groups'][0]['chairman'], '教授A')


class AlgorithmTitleRankTests(TestCase):
    """算法职称等级判断：不允许"教授" in "副教授"这类子串误配"""

    def _teacher(self, title):
        return algorithm.Teacher(id=1, name='教师', title=title)

    def test_associate_professor_does_not_satisfy_professor_requirement(self):
        rules = {'chair_title': '教授'}
        self.assertFalse(algorithm.meets_chair_requirement(self._teacher('副教授'), rules))
        self.assertTrue(algorithm.meets_chair_requirement(self._teacher('教授'), rules))

    def test_higher_title_satisfies_lower_requirement(self):
        rules = {'chair_title': '副教授'}
        self.assertTrue(algorithm.meets_chair_requirement(self._teacher('教授'), rules))
        self.assertFalse(algorithm.meets_chair_requirement(self._teacher('讲师'), rules))

    def test_no_requirement_always_passes(self):
        self.assertTrue(algorithm.meets_chair_requirement(self._teacher('讲师'), {}))
        self.assertTrue(algorithm.meets_chair_requirement(self._teacher(None), {'chair_title': ''}))

    def test_unknown_requirement_falls_back_to_exact_match(self):
        rules = {'chair_title': '特聘研究员'}
        self.assertTrue(algorithm.meets_chair_requirement(self._teacher('特聘研究员'), rules))
        self.assertFalse(algorithm.meets_chair_requirement(self._teacher('教授'), rules))


class ScheduleTypeFilterTests(TestCase):
    """按答辩类型过滤学生与教师资格"""

    BASE_RULES = {
        'defense_type': 'pre',
        'start_date': '2025-05-10',
        'end_date': '2025-05-10',
        'group_size': 2,
        'expert_count': 1,
        'avoid_weekend': False,
        'avoid_supervisor': False,
    }

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_superuser(username='admin', password='strong-pass-123')
        self.client.force_authenticate(self.user)
        Room.objects.create(campus='创新港', name='A101', capacity=30, available_times='')

    def _generate(self, **overrides):
        return self.client.post(
            '/api/schedule/generate/',
            {'rules': {**self.BASE_RULES, **overrides}},
            format='json',
        )

    def _scheduled_student_names(self, response):
        return [
            student['name']
            for group in response.data['groups']
            for student in group['students']
        ]

    def _scheduled_teacher_names(self, response):
        names = []
        for group in response.data['groups']:
            names.extend(teacher['name'] for teacher in group['teachers'])
            names.extend(name for name in [group.get('chairman'), group.get('secretary')] if name)
        return names

    def test_students_filtered_by_defense_type(self):
        Teacher.objects.create(name='老师甲', title='教授')
        Teacher.objects.create(name='老师乙', title='讲师')
        Student.objects.create(
            name='预答辩学生', student_type='学硕', campus='创新港', defense_types=['预答辩'],
        )
        Student.objects.create(
            name='正式答辩学生', student_type='学硕', campus='创新港', defense_types=['正式答辩'],
        )

        response = self._generate()

        self.assertEqual(response.status_code, 200)
        names = self._scheduled_student_names(response)
        self.assertIn('预答辩学生', names)
        self.assertNotIn('正式答辩学生', names)

    def test_student_without_defense_types_skipped_with_notice(self):
        Teacher.objects.create(name='老师甲', title='教授')
        Teacher.objects.create(name='老师乙', title='讲师')
        Student.objects.create(
            name='正常学生', student_type='学硕', campus='创新港', defense_types=['预答辩'],
        )
        Student.objects.create(
            name='漏填学生', student_type='学硕', campus='创新港', defense_types=[],
        )

        response = self._generate()

        self.assertEqual(response.status_code, 200)
        self.assertNotIn('漏填学生', self._scheduled_student_names(response))
        notices = [c for c in response.data['conflicts'] if c['type'] == '数据完整性提示']
        self.assertEqual(len(notices), 1)
        self.assertIn('漏填学生', notices[0]['reason'])
        self.assertEqual(notices[0]['level'], 'warning')

    def test_teachers_filtered_by_available_types(self):
        Teacher.objects.create(name='不限老师', title='教授', available_types=[])
        Teacher.objects.create(name='预答辩老师', title='讲师', available_types=['预答辩'])
        Teacher.objects.create(name='仅正式老师', title='讲师', available_types=['正式答辩'])
        Student.objects.create(
            name='学生1', student_type='学硕', campus='创新港', defense_types=['预答辩'],
        )

        response = self._generate(expert_count=2)

        self.assertEqual(response.status_code, 200)
        names = self._scheduled_teacher_names(response)
        self.assertNotIn('仅正式老师', names)
        self.assertIn('不限老师', names)

    def test_generate_rejects_when_no_students_match_defense_type(self):
        Teacher.objects.create(name='老师甲', title='教授')
        Student.objects.create(
            name='正式答辩学生', student_type='学硕', campus='创新港', defense_types=['正式答辩'],
        )

        response = self._generate()

        self.assertEqual(response.status_code, 400)
        self.assertIn('没有学生参加【预答辩】', response.data['error'])

    def test_generate_rejects_when_no_teachers_match_defense_type(self):
        Teacher.objects.create(name='仅正式老师', title='教授', available_types=['正式答辩'])
        Student.objects.create(
            name='学生1', student_type='学硕', campus='创新港', defense_types=['预答辩'],
        )

        response = self._generate()

        self.assertEqual(response.status_code, 400)
        self.assertIn('没有可参加【预答辩】的教师', response.data['error'])

class AlgorithmGroupSizeTests(TestCase):
    """每组人数上下限的算法行为"""

    @staticmethod
    def _students(count):
        return [algorithm.Student(id=i, name=f'S{i}', campus='创新港') for i in range(1, count + 1)]

    def test_validate_inputs_rejects_end_date_before_start_date(self):
        rules = {
            'start_date': '2025-05-12',
            'end_date': '2025-05-10',
            'group_size': 1,
            'expert_count': 1,
        }

        with self.assertRaisesRegex(algorithm.SchedulingError, 'end_date must be >= start_date'):
            algorithm.validate_inputs(
                teachers=[algorithm.Teacher(id=1, name='教师')],
                students=[algorithm.Student(id=1, name='学生')],
                rooms=[algorithm.Room(id=1, name='A101')],
                rules=rules,
            )

    def test_small_tail_group_merges_into_previous(self):
        rules = {'group_size': 3, 'group_min': 2, 'group_max': 5}
        groups = algorithm.build_student_groups(self._students(4), rules)
        self.assertEqual([len(g) for g in groups], [4])

    def test_tail_group_kept_when_merge_would_exceed_max(self):
        rules = {'group_size': 3, 'group_min': 2, 'group_max': 3}
        groups = algorithm.build_student_groups(self._students(4), rules)
        self.assertEqual([len(g) for g in groups], [3, 1])

    def test_grouping_unchanged_without_min_max(self):
        rules = {'group_size': 3}
        groups = algorithm.build_student_groups(self._students(4), rules)
        self.assertEqual([len(g) for g in groups], [3, 1])

    def test_check_group_size_flags_out_of_range(self):
        rules = {'group_min': 2, 'group_max': 5}
        below = algorithm.check_group_size('G1', 1, rules)
        self.assertEqual(below[0]['type'], 'group_size_out_of_range')
        above = algorithm.check_group_size('G1', 6, rules)
        self.assertEqual(above[0]['type'], 'group_size_out_of_range')
        self.assertEqual(algorithm.check_group_size('G1', 3, rules), [])


class ScheduleRuleEffectTests(TestCase):
    """秘书职称门槛、资深优先、专家下限降级的端到端效果"""

    BASE_RULES = {
        'defense_type': 'pre',
        'start_date': '2025-05-10',
        'end_date': '2025-05-10',
        'group_size': 1,
        'expert_count': 1,
        'avoid_weekend': False,
        'avoid_supervisor': False,
    }

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_superuser(username='admin', password='strong-pass-123')
        self.client.force_authenticate(self.user)
        Room.objects.create(campus='创新港', name='A101', capacity=30, available_times='')

    def _generate(self, **overrides):
        return self.client.post(
            '/api/schedule/generate/',
            {'rules': {**self.BASE_RULES, **overrides}},
            format='json',
        )

    def _setup_three_teachers_one_student(self):
        Teacher.objects.create(name='陈教授', title='教授')
        Teacher.objects.create(name='李副教授', title='副教授')
        Teacher.objects.create(name='张讲师', title='讲师')
        Student.objects.create(
            name='学生1', student_type='学硕', campus='创新港', defense_types=['预答辩'],
        )

    def test_secretary_title_threshold_blocks_low_rank(self):
        # 主席=陈教授、专家=李副教授（资深优先），剩余张讲师低于"副教授"门槛 → 秘书缺位
        self._setup_three_teachers_one_student()

        response = self._generate(
            need_chair=True, chair_title='教授',
            secretary_title='副教授', prefer_senior=True,
        )

        self.assertEqual(response.status_code, 200)
        group = response.data['groups'][0]
        self.assertEqual(group['chairman'], '陈教授')
        self.assertEqual(group['secretary'], '未分配')
        self.assertIn('人员冲突', [c['type'] for c in response.data['conflicts']])

    def test_secretary_title_threshold_allows_qualified(self):
        self._setup_three_teachers_one_student()

        response = self._generate(
            need_chair=True, chair_title='教授',
            secretary_title='讲师', prefer_senior=True,
        )

        self.assertEqual(response.status_code, 200)
        group = response.data['groups'][0]
        self.assertEqual(group['secretary'], '张讲师')

    def test_prefer_senior_picks_highest_title_expert(self):
        self._setup_three_teachers_one_student()

        response = self._generate(prefer_senior=True)

        self.assertEqual(response.status_code, 200)
        expert_names = [t['name'] for t in response.data['groups'][0]['teachers']]
        self.assertEqual(expert_names, ['陈教授'])

    def test_expert_shortage_downgrades_to_warning_when_min_met(self):
        # 导师被回避后可用专家 2 人：低于目标 3 但达到下限 2 → warning 而非 error
        mentor = Teacher.objects.create(name='王教授', title='教授')
        Teacher.objects.create(name='李副教授', title='副教授')
        Teacher.objects.create(name='张讲师', title='讲师')
        Student.objects.create(
            name='学生1', student_type='学硕', mentor_name=mentor.name,
            campus='创新港', defense_types=['预答辩'],
        )

        response = self._generate(avoid_supervisor=True, expert_count=3, expert_min=2)

        self.assertEqual(response.status_code, 200)
        shortage = [c for c in response.data['conflicts'] if 'experts' in c['reason']]
        self.assertEqual(len(shortage), 1)
        self.assertEqual(shortage[0]['level'], 'warning')

    def test_expert_shortage_stays_error_below_min(self):
        mentor = Teacher.objects.create(name='王教授', title='教授')
        Teacher.objects.create(name='李副教授', title='副教授')
        Student.objects.create(
            name='学生1', student_type='学硕', mentor_name=mentor.name,
            campus='创新港', defense_types=['预答辩'],
        )

        response = self._generate(avoid_supervisor=True, expert_count=3, expert_min=2)

        self.assertEqual(response.status_code, 200)
        shortage = [c for c in response.data['conflicts'] if 'experts' in c['reason']]
        self.assertEqual(len(shortage), 1)
        self.assertEqual(shortage[0]['level'], 'error')


class TimeToleranceAndTimezoneTests(TestCase):
    """时间格式容错与本地时区显示"""

    RULES = {
        'defense_type': 'pre',
        'start_date': '2025-05-10',
        'end_date': '2025-05-10',
        'group_size': 1,
        'expert_count': 1,
        'avoid_weekend': False,
        'avoid_supervisor': False,
    }

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_superuser(username='admin', password='strong-pass-123')
        self.client.force_authenticate(self.user)
        Student.objects.create(
            name='学生1', student_type='学硕', campus='创新港', defense_types=['预答辩'],
        )

    def _generate(self):
        return self.client.post('/api/schedule/generate/', {'rules': self.RULES}, format='json')

    def _assigned_teacher_names(self, response):
        names = []
        for group in response.data['groups']:
            names.extend(teacher['name'] for teacher in group['teachers'])
            names.extend(name for name in [group.get('chairman'), group.get('secretary')] if name and name != '未分配')
        return names

    def test_fullwidth_and_slash_time_formats_are_normalized(self):
        # 全角冒号 + 斜杠日期 + 波浪线的不可用时间应被归一化并真实生效
        Teacher.objects.create(
            name='忙碌老师', title='教授',
            unavailable_times='2025/05/10 09：00～12：00',
        )
        Teacher.objects.create(name='空闲老师甲', title='副教授')
        Teacher.objects.create(name='空闲老师乙', title='讲师')
        Room.objects.create(
            campus='创新港', name='A101', capacity=30,
            available_times='2025-05-10 09:00-12:00',
        )

        response = self._generate()

        self.assertEqual(response.status_code, 200)
        self.assertNotIn('忙碌老师', self._assigned_teacher_names(response))
        notices = [c for c in response.data['conflicts'] if c['type'] == '数据完整性提示']
        self.assertEqual(notices, [])

    def test_unparseable_time_generates_notice_instead_of_error(self):
        Teacher.objects.create(name='王老师', title='教授', unavailable_times='周一上午')
        Teacher.objects.create(name='李老师', title='讲师')
        Room.objects.create(campus='创新港', name='A101', capacity=30, available_times='')

        response = self._generate()

        self.assertEqual(response.status_code, 200)
        notices = [c for c in response.data['conflicts'] if c['type'] == '数据完整性提示']
        self.assertEqual(len(notices), 1)
        self.assertIn('王老师', notices[0]['reason'])
        self.assertIn('周一上午', notices[0]['reason'])
        # 无法解析的时间被忽略，教师本身仍参与排期
        self.assertIn('王老师', self._assigned_teacher_names(response))

    def test_room_invalid_time_falls_back_with_notice(self):
        Teacher.objects.create(name='王老师', title='教授')
        Teacher.objects.create(name='李老师', title='讲师')
        Room.objects.create(campus='创新港', name='B202', capacity=30, available_times='随时')

        response = self._generate()

        self.assertEqual(response.status_code, 200)
        notices = [c for c in response.data['conflicts'] if c['type'] == '数据完整性提示']
        self.assertEqual(len(notices), 1)
        self.assertIn('B202', notices[0]['reason'])
        # 教室时间被忽略后回退到默认时段，分组仍可排出时间
        self.assertTrue(response.data['groups'][0]['date'])

    def test_generated_at_uses_local_timezone(self):
        from django.utils import timezone as dj_timezone

        Teacher.objects.create(name='王老师', title='教授')
        Teacher.objects.create(name='李老师', title='讲师')
        Room.objects.create(campus='创新港', name='A101', capacity=30, available_times='')

        response = self._generate()

        self.assertEqual(response.status_code, 200)
        version = ScheduleVersion.objects.get(defense_type='pre', is_current=True)
        expected = dj_timezone.localtime(version.created_at).strftime('%Y-%m-%d %H:%M')
        self.assertEqual(response.data['generatedAt'], expected)


class TimetableImportTests(TestCase):
    """课表导入：矩阵式/行式课表 → 教师不可用时间"""

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_superuser(username='admin', password='strong-pass-123')
        self.client.force_authenticate(self.user)

    def _matrix_upload(self):
        # 模拟真实矩阵课表的脏格式：全角/半角/无括号周次、双师、周次与教师名分行、连写姓名
        return make_xlsx_upload(
            '软件学院课表.xlsx',
            [
                ['2025—2026学年第二学期课程表', None, None, None, None, None],
                ['上课时间', None, '星期一', None, '星期二', None],
                [None, None, '1-8周', '9-16周', '1-8周', '9-16周'],
                ['上午', '1节', '【5-12】\n张甲', None, '【9-16】', None],
                [None, '2节', None, None, '李乙', None],
                [None, '3节', '【1-8】王丙 赵丁', '9-16\n钱戊', None, '[9-16]\n孙己'],
                [None, '4节', None, None, None, '郑帅祝继华'],
            ],
        )

    def test_matrix_timetable_import_expands_teacher_busy_times(self):
        for name in ['张甲', '李乙', '王丙', '赵丁', '钱戊', '孙己', '郑帅', '祝继华']:
            Teacher.objects.create(name=name, title='副教授')

        response = self.client.post(
            '/api/teachers/import_timetable/',
            {'file': self._matrix_upload(), 'semesterFirstMonday': '2026-03-02'},
            format='multipart',
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['unknownTeachers'], 0)

        zhang = Teacher.objects.get(name='张甲')
        zhang_entries = zhang.unavailable_times.split('\n')
        # 5-12 周共 8 周，周一 1-2 节 08:00-10:00，第 5 周周一 = 2026-03-30
        self.assertEqual(len(zhang_entries), 8)
        self.assertIn('2026-03-30 08:00-10:00', zhang_entries)

        # 周次【9-16】与教师名"李乙"分处上下两行单元格（真实课表常见溢出）
        li = Teacher.objects.get(name='李乙')
        self.assertEqual(len(li.unavailable_times.split('\n')), 8)
        self.assertIn('2026-04-28 08:00-10:00', li.unavailable_times)

        # 同格双师
        for name in ['王丙', '赵丁']:
            teacher = Teacher.objects.get(name=name)
            self.assertIn('2026-03-02 10:00-12:00', teacher.unavailable_times)

        # 无括号周次 "9-16\n钱戊"
        qian = Teacher.objects.get(name='钱戊')
        self.assertEqual(len(qian.unavailable_times.split('\n')), 8)

        # 半角括号 "[9-16]\n孙己"
        sun = Teacher.objects.get(name='孙己')
        self.assertEqual(len(sun.unavailable_times.split('\n')), 8)

        # 连写姓名"郑帅祝继华"按系统教师名单切分（星期二 3-4 节，列默认 9-16 周）
        zheng = Teacher.objects.get(name='郑帅')
        zhu = Teacher.objects.get(name='祝继华')
        self.assertIn('2026-04-28 10:00-12:00', zheng.unavailable_times)
        self.assertIn('2026-04-28 10:00-12:00', zhu.unavailable_times)

    def test_roster_timetable_import_parses_time_location_text(self):
        Teacher.objects.create(name='高庚', title='教授')
        Teacher.objects.create(name='周辛', title='讲师')

        upload = make_xlsx_upload(
            '学部课表.xlsx',
            [
                ['课程名称', '主讲教师', '时间地点'],
                ['最优控制', '高庚', '创新港校区-5-6043周次:第9-16周 连续周 星期三 上3,上4,下7,下8'],
                ['网络', '周辛', '兴庆校区-中2周次:第1-2周 连续周 星期二 上1,上2 星期四 下5,下6'],
            ],
        )

        response = self.client.post(
            '/api/teachers/import_timetable/',
            {'file': upload, 'semesterFirstMonday': '2026-03-02'},
            format='multipart',
        )

        self.assertEqual(response.status_code, 200)
        gao = Teacher.objects.get(name='高庚')
        gao_entries = gao.unavailable_times.split('\n')
        # 8 周 × (3-4 节 + 7-8 节) = 16 条；第 9 周周三 = 2026-04-29
        self.assertEqual(len(gao_entries), 16)
        self.assertIn('2026-04-29 10:00-12:00', gao_entries)
        self.assertIn('2026-04-29 16:00-18:00', gao_entries)

        zhou = Teacher.objects.get(name='周辛')
        zhou_entries = zhou.unavailable_times.split('\n')
        # 2 周 × 2 段 = 4 条：周二上午 + 周四下午
        self.assertEqual(len(zhou_entries), 4)
        self.assertIn('2026-03-03 08:00-10:00', zhou_entries)
        self.assertIn('2026-03-05 14:00-16:00', zhou_entries)

    def test_timetable_import_is_idempotent(self):
        Teacher.objects.create(name='张甲', title='副教授')

        for _ in range(2):
            response = self.client.post(
                '/api/teachers/import_timetable/',
                {'file': self._matrix_upload(), 'semesterFirstMonday': '2026-03-02'},
                format='multipart',
            )
            self.assertEqual(response.status_code, 200)

        zhang = Teacher.objects.get(name='张甲')
        self.assertEqual(len(zhang.unavailable_times.split('\n')), 8)

    def test_timetable_import_skips_unknown_teachers_and_keeps_existing_times(self):
        Teacher.objects.create(
            name='张甲', title='副教授', unavailable_times='2026-05-01 09:00-12:00'
        )

        response = self.client.post(
            '/api/teachers/import_timetable/',
            {'file': self._matrix_upload(), 'semesterFirstMonday': '2026-03-02'},
            format='multipart',
        )

        self.assertEqual(response.status_code, 200)
        self.assertGreater(response.data['unknownTeachers'], 0)
        self.assertNotIn('张甲', response.data['unknownTeacherNames'])

        zhang = Teacher.objects.get(name='张甲')
        entries = zhang.unavailable_times.split('\n')
        # 手工录入的既有条目保留，课表条目追加
        self.assertEqual(entries[0], '2026-05-01 09:00-12:00')
        self.assertEqual(len(entries), 9)

    def test_timetable_import_rejects_non_monday_semester_start(self):
        response = self.client.post(
            '/api/teachers/import_timetable/',
            {'file': self._matrix_upload(), 'semesterFirstMonday': '2026-03-03'},
            format='multipart',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('周一', response.data['error'])

    def test_timetable_import_requires_semester_start_date(self):
        response = self.client.post(
            '/api/teachers/import_timetable/',
            {'file': self._matrix_upload()},
            format='multipart',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('semesterFirstMonday', response.data['error'])


REAL_DOCS_DIR = Path(__file__).resolve().parent.parent / '排答辩文档'


def real_doc_upload(filename):
    path = REAL_DOCS_DIR / filename
    return SimpleUploadedFile(filename, path.read_bytes())


@unittest.skipUnless(REAL_DOCS_DIR.exists(), '真实答辩文档目录不存在，跳过冒烟验证')
class RealDocumentSmokeTests(TestCase):
    """用 排答辩文档/ 下的真实文件做端到端导入冒烟验证（仅在本机文件存在时运行）"""

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_superuser(username='admin', password='strong-pass-123')
        self.client.force_authenticate(self.user)

    def test_all_real_documents_import_end_to_end(self):
        # 1. 教师：指定组长秘书（角色列无表头，靠内容嗅探）
        response = self.client.post(
            '/api/teachers/import_data/',
            {'file': real_doc_upload('指定组长秘书.xlsx')},
            format='multipart',
        )
        self.assertEqual(response.status_code, 200, response.data)
        teacher_count = Teacher.objects.count()
        self.assertGreater(teacher_count, 30)
        self.assertTrue(Teacher.objects.exclude(roles=[]).exists())

        # 2. 学生：23 级答辩地点统计（标题行 + 学号/性别/学科/导师）
        response = self.client.post(
            '/api/students/import_data/',
            {'file': real_doc_upload('23级答辩地点统计-含学生信息-导师信息-答辩地点.xlsx')},
            format='multipart',
        )
        self.assertEqual(response.status_code, 200, response.data)
        formal_students = Student.objects.count()
        self.assertGreater(formal_students, 300)
        # 学号全部落库且非空
        self.assertEqual(Student.objects.filter(student_no__isnull=True).count(), 0)
        self.assertGreater(Student.objects.exclude(gender='').count(), 300)

        # 3. 学生：22 级中期考核答辩地点（空行后表头，另一批学生）
        response = self.client.post(
            '/api/students/import_data/',
            {'file': real_doc_upload('22级中期考核答辩地点.xlsx')},
            format='multipart',
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertGreater(Student.objects.count(), formal_students)
        # SQLite 的 JSONField 不支持 contains 查询，用 Python 侧断言
        self.assertTrue(any(
            '中期答辩' in (student.defense_types or [])
            for student in Student.objects.all()
        ))

        # 4. 教室：创新港（卡片式导出格式，.xlsx）
        response = self.client.post(
            '/api/rooms/import_data/',
            {'file': real_doc_upload('教室借用申请-创新港教室.xlsx')},
            format='multipart',
        )
        self.assertEqual(response.status_code, 200, response.data)
        innovation_rooms = Room.objects.filter(campus='创新港').count()
        self.assertGreater(innovation_rooms, 0)

        # 5. 教室：兴庆（教务系统标准表格，.xls 老格式）
        response = self.client.post(
            '/api/rooms/import_data/',
            {'file': real_doc_upload('教室借用申请-兴庆教室.xls')},
            format='multipart',
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertGreater(Room.objects.filter(campus='兴庆').count(), 0)

        # 6. 课表：软件学院矩阵式课表（.xlsx）
        response = self.client.post(
            '/api/teachers/import_timetable/',
            {
                'file': real_doc_upload('2025—2026学年第二学期软件学院课表(1).xlsx'),
                'semesterFirstMonday': '2026-03-02',
            },
            format='multipart',
        )
        self.assertEqual(response.status_code, 200, response.data)
        matrix_updated = response.data['updatedTeachers']
        self.assertGreater(matrix_updated, 0)
        self.assertEqual(response.data['warnings'], [])

        # 7. 课表：学部行式课表（.xls）
        response = self.client.post(
            '/api/teachers/import_timetable/',
            {
                'file': real_doc_upload('学部课表.xls'),
                'semesterFirstMonday': '2026-03-02',
            },
            format='multipart',
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['warnings'], [])

        # 课表导入后确有教师拿到了格式合法的不可用时间
        busy_teacher = Teacher.objects.exclude(unavailable_times='').first()
        self.assertIsNotNone(busy_teacher)
        from .views import parse_time_entries
        valid_entries, invalid_entries = parse_time_entries(busy_teacher.unavailable_times)
        self.assertGreater(len(valid_entries), 0)
        self.assertEqual(invalid_entries, [])

    def _import_real_documents(self):
        """依次导入全部真实文件（每步断言成功），供全链路演练复用"""
        steps = [
            ('/api/teachers/import_data/', '指定组长秘书.xlsx', {}),
            ('/api/students/import_data/', '23级答辩地点统计-含学生信息-导师信息-答辩地点.xlsx', {}),
            ('/api/students/import_data/', '22级中期考核答辩地点.xlsx', {}),
            ('/api/rooms/import_data/', '教室借用申请-创新港教室.xlsx', {}),
            ('/api/rooms/import_data/', '教室借用申请-兴庆教室.xls', {}),
            ('/api/teachers/import_timetable/', '2025—2026学年第二学期软件学院课表(1).xlsx',
             {'semesterFirstMonday': '2026-03-02'}),
            ('/api/teachers/import_timetable/', '学部课表.xls', {'semesterFirstMonday': '2026-03-02'}),
        ]
        for url, filename, extra in steps:
            payload = {'file': real_doc_upload(filename), **extra}
            response = self.client.post(url, payload, format='multipart')
            self.assertEqual(response.status_code, 200, f'{filename}: {getattr(response, "data", "")}')

    def _append_room_slots(self, start_day, end_day):
        """给教室补录指定日期范围的工作日可用时段（模拟新学期教室借用到位）"""
        slots = []
        current = start_day
        while current <= end_day:
            if current.weekday() < 5:
                slots.append(f'{current.isoformat()} 09:00-12:00')
                slots.append(f'{current.isoformat()} 14:00-18:00')
            current += timedelta(days=1)
        for room in Room.objects.all():
            entries = [entry for entry in (room.available_times or '').split('\n') if entry]
            room.available_times = '\n'.join(entries + slots)
            room.save(update_fields=['available_times'])

    def test_full_real_data_schedule_pipeline(self):
        self._import_real_documents()

        # 23 级学生先参加预答辩、再参加正式答辩（同一批人，正式沿用预答辩分组）
        formal_students = [s for s in Student.objects.all() if '正式答辩' in (s.defense_types or [])]
        self.assertGreater(len(formal_students), 300)
        for student in formal_students:
            student.defense_types = ['预答辩'] + [t for t in student.defense_types if t != '预答辩']
            student.save(update_fields=['defense_types'])

        # 真实借用记录只有 4 月（预答辩）教室：给 5 月正式答辩、7 月中期补录教室时段
        self._append_room_slots(date(2026, 5, 6), date(2026, 5, 22))
        self._append_room_slots(date(2026, 7, 6), date(2026, 7, 17))

        base_rules = {
            'group_size': 6, 'group_min': 4, 'group_max': 8,
            'expert_count': 3, 'expert_min': 3,
            'need_chair': True, 'chair_title': '副教授',
            'avoid_weekend': True,
        }

        # 1) 预答辩：导师必须同组，按导师聚类分组，避开清明假期
        response = self.client.post('/api/schedule/generate/', {'rules': {
            **base_rules,
            'defense_type': 'pre',
            'start_date': '2026-04-01', 'end_date': '2026-04-30',
            'exclude_dates': ['2026-04-06'],
            'supervisor_policy': 'same_group', 'grouping': 'supervisor',
        }}, format='json')
        self.assertEqual(response.status_code, 200, getattr(response, 'data', ''))

        pre_version = ScheduleVersion.objects.get(defense_type='pre', is_current=True)
        pre_group_count = pre_version.groups.count()
        self.assertGreater(pre_group_count, 40)

        # 预答辩生成后秘书绑定已回写学生档案
        bound_students = Student.objects.exclude(secretary_name='').count()
        self.assertGreater(bound_students, 250)

        # 清明及周末未被排期
        for group in pre_version.groups.all():
            if group.time:
                day = date.fromisoformat(group.time.split(' ')[0])
                self.assertNotEqual(day.isoformat(), '2026-04-06')
                self.assertLess(day.weekday(), 5)

        # 2) 正式答辩：导师回避，按秘书聚类沿用预答辩分组
        response = self.client.post('/api/schedule/generate/', {'rules': {
            **base_rules,
            'defense_type': 'formal',
            'start_date': '2026-05-06', 'end_date': '2026-05-22',
            'expert_count': 4, 'expert_min': 4, 'chair_title': '教授',
            'supervisor_policy': 'avoid', 'grouping': 'secretary',
        }}, format='json')
        self.assertEqual(response.status_code, 200, getattr(response, 'data', ''))

        formal_version = ScheduleVersion.objects.get(defense_type='formal', is_current=True)
        pre_sets = {
            frozenset(group.students.values_list('id', flat=True))
            for group in pre_version.groups.all()
        }
        formal_sets = {
            frozenset(group.students.values_list('id', flat=True))
            for group in formal_version.groups.all()
        }
        preserved = len(pre_sets & formal_sets)
        # 绝大多数组从预答辩原样延续（个别组因秘书缺失可能被重新装箱）
        self.assertGreater(preserved / max(len(pre_sets), 1), 0.7,
                           f'仅 {preserved}/{len(pre_sets)} 组保持不变')

        # 3) 中期考核：22 级学生，每组 10-12 人、专家 5 人（含组长）
        response = self.client.post('/api/schedule/generate/', {'rules': {
            'defense_type': 'mid',
            'start_date': '2026-07-06', 'end_date': '2026-07-17',
            'group_size': 11, 'group_min': 10, 'group_max': 12,
            'expert_count': 4, 'expert_min': 4,
            'need_chair': True, 'chair_title': '副教授',
            'avoid_weekend': True,
            'supervisor_policy': 'same_group', 'grouping': 'supervisor',
        }}, format='json')
        self.assertEqual(response.status_code, 200, getattr(response, 'data', ''))

        # 4) 三种答辩类型的 Word 与 Excel 导出全部可用
        for defense_type in ('pre', 'formal', 'mid'):
            word_response = self.client.get('/api/schedule/export_word/', {'defense_type': defense_type})
            self.assertEqual(word_response.status_code, 200, defense_type)
            self.assertIn('wordprocessingml', word_response['Content-Type'])

            excel_response = self.client.get('/api/schedule/export/', {'defense_type': defense_type})
            self.assertEqual(excel_response.status_code, 200, defense_type)


class SupervisorPolicyTests(TestCase):
    """导师约束三态、秘书跟随绑定与聚类分组的算法行为"""

    def _teachers(self):
        return [
            {'id': 1, 'name': '导师一', 'title': '副教授'},
            {'id': 2, 'name': '导师二', 'title': '副教授'},
            {'id': 3, 'name': '导师三', 'title': '副教授'},
            {'id': 4, 'name': '主席候选', 'title': '教授'},
            {'id': 5, 'name': '秘书候选', 'title': '讲师'},
            {'id': 6, 'name': '专家甲', 'title': '副教授'},
            {'id': 7, 'name': '专家乙', 'title': '副教授'},
            {'id': 8, 'name': '专家丙', 'title': '副教授'},
        ]

    def _students(self, secretary_id=None):
        supervisors = [1, 1, 2, 2, 3, 3]
        return [
            {
                'id': 100 + index,
                'name': f'学生{index}',
                'supervisor_id': supervisors[index - 1],
                'campus': '创新港',
                'secretary_id': secretary_id,
            }
            for index in range(1, 7)
        ]

    def _rules(self, **overrides):
        rules = {
            'defense_type': 'pre',
            'start_date': '2026-04-06',
            'end_date': '2026-04-10',
            'group_size': 6,
            'group_min': 3,
            'group_max': 8,
            'expert_count': 3,
            'expert_min': 3,
            'need_chair': True,
            'chair_title': '教授',
            'prefer_senior': True,
            'avoid_weekend': True,
        }
        rules.update(overrides)
        return rules

    def _generate(self, teachers=None, students=None, rooms=None, **rule_overrides):
        return algorithm.generate_schedule(
            teachers=teachers or self._teachers(),
            students=students or self._students(),
            rooms=rooms or [{'id': 1, 'campus': '创新港', 'name': 'A101'}],
            rules=self._rules(**rule_overrides),
        )

    def test_same_group_policy_places_supervisors_with_students(self):
        result = self._generate(supervisor_policy='same_group')

        self.assertEqual(len(result['groups']), 1)
        group = result['groups'][0]
        assigned = {group['chair_id'], group['secretary_id'], *group['expert_ids']}
        for supervisor_id in (1, 2, 3):
            self.assertIn(supervisor_id, assigned)
        conflict_types = {c['type'] for c in result['conflicts']}
        self.assertNotIn('supervisor_missing', conflict_types)
        # 秘书不能是组内学生的导师
        self.assertNotIn(group['secretary_id'], {1, 2, 3})

    def test_avoid_policy_keeps_supervisors_out(self):
        result = self._generate(supervisor_policy='avoid')

        group = result['groups'][0]
        assigned = {group['chair_id'], group['secretary_id'], *group['expert_ids']}
        for supervisor_id in (1, 2, 3):
            self.assertNotIn(supervisor_id, assigned)
        conflict_types = {c['type'] for c in result['conflicts']}
        self.assertNotIn('supervisor_avoidance', conflict_types)

    def test_secretary_follows_students_previous_binding(self):
        result = self._generate(
            students=self._students(secretary_id=5),
            supervisor_policy='avoid',
            grouping='secretary',
        )

        group = result['groups'][0]
        self.assertEqual(group['secretary_id'], 5)
        conflict_types = {c['type'] for c in result['conflicts']}
        self.assertNotIn('secretary_continuity_broken', conflict_types)

    def test_bound_secretary_rejected_when_supervising_group_students(self):
        # 学生绑定的秘书同时是组内两名学生的导师：绑定被拒并生成连续性提示
        students = self._students(secretary_id=5)
        students[0]['supervisor_id'] = 5
        students[1]['supervisor_id'] = 5

        result = self._generate(students=students, supervisor_policy='none')

        group = result['groups'][0]
        self.assertNotEqual(group['secretary_id'], 5)
        conflict_types = {c['type'] for c in result['conflicts']}
        self.assertIn('secretary_continuity_broken', conflict_types)

    def test_supervisor_grouping_keeps_same_mentor_students_together(self):
        supervisors = [1] * 4 + [2] * 4 + [3] * 4
        students = [
            {'id': 200 + index, 'name': f'学生{index}', 'supervisor_id': supervisors[index - 1], 'campus': '创新港'}
            for index in range(1, 13)
        ]

        result = self._generate(students=students, supervisor_policy='same_group', grouping='supervisor')

        student_map = {s['id']: s for s in students}
        for group in result['groups']:
            group_supervisors = {student_map[sid]['supervisor_id'] for sid in group['student_ids']}
            # 每组内不应出现"同导师学生被拆到别组"的情况
            for supervisor_id in group_supervisors:
                all_ids = {s['id'] for s in students if s['supervisor_id'] == supervisor_id}
                self.assertTrue(all_ids.issubset(set(group['student_ids'])))

    def test_secretary_grouping_reuses_previous_groups(self):
        students = []
        for index in range(1, 13):
            students.append({
                'id': 300 + index,
                'name': f'学生{index}',
                'supervisor_id': 1 if index <= 6 else 2,
                'campus': '创新港',
                'secretary_id': 5 if index <= 6 else 6,
            })

        result = self._generate(
            students=students,
            supervisor_policy='avoid',
            grouping='secretary',
            rooms=[
                {'id': 1, 'campus': '创新港', 'name': 'A101'},
                {'id': 2, 'campus': '创新港', 'name': 'A102'},
            ],
        )

        self.assertEqual(len(result['groups']), 2)
        grouped_ids = [set(group['student_ids']) for group in result['groups']]
        self.assertIn({301, 302, 303, 304, 305, 306}, grouped_ids)
        self.assertIn({307, 308, 309, 310, 311, 312}, grouped_ids)

    def test_exclude_dates_skips_configured_holidays(self):
        result = self._generate(
            supervisor_policy='same_group',
            start_date='2026-04-06',
            end_date='2026-04-07',
            exclude_dates=['2026-04-06'],
        )

        group = result['groups'][0]
        self.assertIsNotNone(group['time'])
        self.assertTrue(group['time'].startswith('2026-04-07'), group['time'])

    def test_exclude_dates_filters_room_available_slots(self):
        rooms = [{
            'id': 1, 'campus': '创新港', 'name': 'A101',
            'available_time': ['2026-04-06 09:00-12:00', '2026-04-07 09:00-12:00'],
        }]

        result = self._generate(
            rooms=rooms,
            supervisor_policy='same_group',
            exclude_dates=['2026-04-06'],
        )

        group = result['groups'][0]
        self.assertTrue((group['time'] or '').startswith('2026-04-07'), group['time'])


class SoftWeightTests(TestCase):
    """软约束权重对分组与教师挑选的影响"""

    def _rules(self, **overrides):
        rules = {
            'defense_type': 'pre',
            'start_date': '2026-04-06',
            'end_date': '2026-04-10',
            'group_size': 4,
            'group_min': 2,
            'group_max': 6,
            'expert_count': 1,
            'need_chair': False,
            'avoid_weekend': True,
            'supervisor_policy': 'none',
            'grouping': 'supervisor',
        }
        rules.update(overrides)
        return rules

    def test_low_balance_weight_fills_groups_to_max(self):
        # 6 位导师各带 2 名学生：均衡权重高 → 贴目标人数 4 人/组；权重低 → 填到上限 6 人/组
        students = [
            {'id': 400 + index, 'name': f'学生{index}', 'supervisor_id': (index + 1) // 2, 'campus': '创新港'}
            for index in range(1, 13)
        ]
        teachers = [{'id': i, 'name': f'教师{i}', 'title': '副教授'} for i in range(1, 10)]
        rooms = [{'id': i, 'campus': '创新港', 'name': f'A{i}'} for i in range(1, 5)]

        balanced = algorithm.build_student_groups(
            [algorithm.parse_student(s) for s in students],
            self._rules(soft_weights={'balance_student_count': 80}),
        )
        compact = algorithm.build_student_groups(
            [algorithm.parse_student(s) for s in students],
            self._rules(soft_weights={'balance_student_count': 20}),
        )

        self.assertEqual(sorted(len(g) for g in balanced), [4, 4, 4])
        self.assertEqual(sorted(len(g) for g in compact), [6, 6])

    def test_cross_campus_weight_defers_teachers_already_elsewhere(self):
        teachers = [
            {'id': 1, 'name': '教师甲', 'title': '副教授'},
            {'id': 2, 'name': '教师乙', 'title': '副教授'},
            {'id': 3, 'name': '教师丙', 'title': '副教授'},
            {'id': 4, 'name': '教师丁', 'title': '副教授'},
        ]
        students = [
            {'id': 501, 'name': '港学生', 'campus': '创新港'},
            {'id': 502, 'name': '兴学生', 'campus': '兴庆'},
        ]
        rooms = [
            {'id': 1, 'campus': '创新港', 'name': 'A101', 'available_time': ['2026-04-07 09:00-12:00']},
            {'id': 2, 'campus': '兴庆', 'name': 'B201', 'available_time': ['2026-04-07 14:00-17:00']},
        ]

        def run(weight):
            return algorithm.generate_schedule(
                teachers=teachers,
                students=students,
                rooms=rooms,
                rules=self._rules(
                    group_size=1, group_min=1, group_max=1,
                    soft_weights={'avoid_cross_campus': weight},
                ),
            )

        without_penalty = run(0)
        with_penalty = run(80)

        # 兴庆组先处理并占用教师甲乙；后处理的创新港组在无惩罚时继续用教师甲，
        # 有惩罚时甲（当天已在兴庆）被排到队尾，换成当天空闲的其他教师
        second_group_without = next(g for g in without_penalty['groups'] if g['campus'] == '创新港')
        second_group_with = next(g for g in with_penalty['groups'] if g['campus'] == '创新港')
        self.assertIn(1, second_group_without['expert_ids'])
        self.assertNotIn(1, second_group_with['expert_ids'])

    def test_academic_master_weight_gives_master_groups_earlier_slots(self):
        students = (
            [{'id': 600 + i, 'name': f'专硕{i}', 'type': '专硕', 'supervisor_id': 1, 'campus': '创新港'} for i in range(1, 5)]
            + [{'id': 610 + i, 'name': f'学硕{i}', 'type': '学硕', 'supervisor_id': 2, 'campus': '创新港'} for i in range(1, 5)]
        )
        teachers = [{'id': i, 'name': f'教师{i}', 'title': '副教授'} for i in range(1, 8)]
        rooms = [{'id': 1, 'campus': '创新港', 'name': 'A101'}]

        result = algorithm.generate_schedule(
            teachers=teachers,
            students=students,
            rooms=rooms,
            rules=self._rules(soft_weights={'prefer_academic_master_first': 80}),
        )

        groups_by_time = sorted(result['groups'], key=lambda g: g['time'] or '9999')
        first_students = set(groups_by_time[0]['student_ids'])
        self.assertTrue(first_students.issuperset({611, 612, 613, 614}))


class ScheduleExportColorTests(TestCase):
    """Word/Excel 导出的导师-学生同色标注"""

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_superuser(username='admin', password='strong-pass-123')
        self.client.force_authenticate(self.user)

        Teacher.objects.create(name='导师王', title='副教授')
        Teacher.objects.create(name='导师李', title='副教授')
        Teacher.objects.create(name='主席张', title='教授')
        Teacher.objects.create(name='秘书赵', title='讲师')
        Teacher.objects.create(name='专家钱', title='副教授')
        for index in range(1, 7):
            Student.objects.create(
                name=f'学生{index}',
                student_no=f'31231580{index:02d}',
                mentor_name='导师王' if index <= 3 else '导师李',
                campus='创新港',
                defense_types=['预答辩'],
            )
        Room.objects.create(campus='创新港', name='A101', capacity=30)

    def _generate(self):
        return self.client.post('/api/schedule/generate/', {
            'rules': {
                'defense_type': 'pre',
                'start_date': '2026-04-06',
                'end_date': '2026-04-10',
                'group_size': 6,
                'group_min': 3,
                'group_max': 8,
                'expert_count': 2,
                'expert_min': 2,
                'need_chair': True,
                'chair_title': '教授',
                'supervisor_policy': 'same_group',
                'grouping': 'supervisor',
                'avoid_weekend': True,
            },
        }, format='json')

    def test_word_export_colors_students_with_their_mentors(self):
        self.assertEqual(self._generate().status_code, 200)

        response = self.client.get('/api/schedule/export_word/', {'defense_type': 'pre'})

        self.assertEqual(response.status_code, 200)
        self.assertIn('wordprocessingml', response['Content-Type'])

        from docx import Document as ReadDocument
        doc = ReadDocument(BytesIO(response.content))
        self.assertIn('预答辩时间安排', doc.paragraphs[0].text)

        run_colors = {}
        for table in doc.tables:
            for docx_row in table.rows:
                for cell in docx_row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            text = run.text.strip()
                            if text and run.font.color and run.font.color.rgb:
                                run_colors[text] = str(run.font.color.rgb)

        # 结构标签为红色
        self.assertEqual(run_colors.get('时间'), 'FF0000')
        # 学生与其导师（组内专家）同色且非红
        self.assertIsNotNone(run_colors.get('学生1'))
        self.assertEqual(run_colors.get('学生1'), run_colors.get('导师王'))
        self.assertEqual(run_colors.get('学生4'), run_colors.get('导师李'))
        self.assertNotEqual(run_colors.get('学生1'), run_colors.get('学生4'))
        self.assertNotEqual(run_colors.get('学生1'), 'FF0000')

    def test_excel_export_uses_mentor_font_colors(self):
        self.assertEqual(self._generate().status_code, 200)

        response = self.client.get('/api/schedule/export/', {'defense_type': 'pre'})

        self.assertEqual(response.status_code, 200)

        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(response.content), rich_text=True)
        sheet = workbook.worksheets[0]

        color_by_text = {}
        for excel_row in sheet.iter_rows():
            for cell in excel_row:
                if isinstance(cell.value, str) and cell.font and cell.font.color and cell.font.color.rgb:
                    color_by_text[cell.value] = str(cell.font.color.rgb)

        student_color = color_by_text.get('学生1')
        mentor_color = color_by_text.get('导师王')
        self.assertIsNotNone(student_color)
        self.assertEqual(student_color, mentor_color)
        other_color = color_by_text.get('学生4')
        self.assertIsNotNone(other_color)
        self.assertNotEqual(student_color, other_color)

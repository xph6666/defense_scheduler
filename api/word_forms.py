"""Word 导出共用的排版基元和字段文案。

`答辩时间安排`（`api/export_word.py`）与 `评审小组结论`（`api/export_conclusion.py`）
是学院归档的两份固定格式文档，版式数值全部取自真实归档样例：

- `test-data/26年数据/学位论文答辩/预答辩/软件工程硕士研究生2026年4月学位论文预答辩时间安排-修订稿.docx`
- `test-data/26年数据/学位论文答辩/预答辩/软件工程硕士研究生学位论文预答辩评审小组结论-2026.docx`

本模块只放两份文档都需要的底层 OOXML 设置（中文字体、表格边框与固定列宽、行高、
段落间距）以及时间段文案格式，具体版式仍留在各自的导出模块里。

中文字体必须同时设置 ``w:ascii`` 与 ``w:eastAsia`` 并带 ``w:hint="eastAsia"``，否则
Word 会用西文字体回退渲染汉字，归档文档会明显走形。
"""
import re
from datetime import date

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

# 学院归档样例使用的字体。华文行楷等字体缺失时 Word 自行回退，不影响打开。
SONG = '宋体'
HEI = '黑体'
FANG_SONG = '仿宋_GB2312'
KAI = '楷体_GB2312'
XING_KAI = '华文行楷'

BLACK = '000000'
RED = 'FF0000'
ERROR = 'C00000'
LIGHT_GRAY = 'C0C0C0'

ALIGNMENTS = {
    'left': WD_ALIGN_PARAGRAPH.LEFT,
    'center': WD_ALIGN_PARAGRAPH.CENTER,
    'right': WD_ALIGN_PARAGRAPH.RIGHT,
    'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
}

WEEKDAYS = '一二三四五六日'

# tblPr 子元素的 schema 顺序：tblW 与 tblInd/tblBorders 必须插在它们该在的位置，
# Word 对顺序敏感，插错会让整个表格设置不生效。
_AFTER_TBL_W = ('w:jc', 'w:tblCellSpacing', 'w:tblInd', 'w:tblBorders', 'w:shd',
                'w:tblLayout', 'w:tblCellMar', 'w:tblLook', 'w:tblCaption',
                'w:tblDescription', 'w:tblPrChange')
_AFTER_TBL_IND = ('w:tblBorders', 'w:shd', 'w:tblLayout', 'w:tblCellMar',
                  'w:tblLook', 'w:tblCaption', 'w:tblDescription', 'w:tblPrChange')
_AFTER_TBL_BORDERS = ('w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook',
                      'w:tblCaption', 'w:tblDescription', 'w:tblPrChange')


# ---------------------------------------------------------------- 文档与段落 ----

def configure_document(doc, *, width_cm, height_cm, left_cm, right_cm, top_cm, bottom_cm,
                       base_font=SONG, base_size=10.5):
    """设置页面尺寸、页边距和 Normal 样式的默认中文字体。

    页边距与纸张必须逐项设置：python-docx 默认纸型是 Letter，直接沿用会让学院的
    A4 表单被挤成两页。
    """
    section = doc.sections[0]
    section.page_width = Cm(width_cm)
    section.page_height = Cm(height_cm)
    section.left_margin = Cm(left_cm)
    section.right_margin = Cm(right_cm)
    section.top_margin = Cm(top_cm)
    section.bottom_margin = Cm(bottom_cm)
    style = doc.styles['Normal']
    style.font.name = base_font
    style.font.size = Pt(base_size)
    # python-docx 的默认模板给所有段落留 10 磅「段后」，表格单元格也继承，
    # 归档表里的多行单元格会因此被撑到下一页；这里统一清零并按单倍行距排版。
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.0
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn('w:eastAsia'), base_font)
    r_fonts.set(qn('w:hint'), 'eastAsia')
    return section


def styled_run(paragraph, text, *, font=SONG, size=10.5, bold=False, color=None,
               underline=False, superscript=False):
    """写入一个带中文字体的 run；颜色为空时保持黑色。"""
    run = paragraph.add_run(text)
    run.font.name = font
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn('w:eastAsia'), font)
    r_fonts.set(qn('w:hint'), 'eastAsia')
    if size:
        run.font.size = Pt(size)
    run.bold = bool(bold)
    if color and color != BLACK:
        run.font.color.rgb = RGBColor.from_string(color)
    if underline:
        run.font.underline = True
    if superscript:
        run.font.superscript = True
    return run


def add_paragraph(container, *, align=None, line=None, line_rule='auto', before=None,
                  after=None, right_indent=None, first_line_indent=None,
                  page_break_before=False):
    """新建段落并按归档样例设置对齐、行距和缩进。

    ``line`` 以缇（1/20 磅）为单位：``line=480`` 配 ``line_rule='auto'`` 是 2 倍行距，
    配 ``'exact'`` 是固定 24 磅。归档文档大量使用两种混排，必须分开表达。
    """
    return format_paragraph(container.add_paragraph(), align=align, line=line,
                            line_rule=line_rule, before=before, after=after,
                            right_indent=right_indent, first_line_indent=first_line_indent,
                            page_break_before=page_break_before)


def format_paragraph(paragraph, *, align=None, line=None, line_rule='auto', before=None,
                     after=None, right_indent=None, first_line_indent=None,
                     page_break_before=False):
    """设置已有段落（例如单元格自带的空段）的对齐、行距和缩进。"""
    if align:
        paragraph.alignment = ALIGNMENTS[align]
    fmt = paragraph.paragraph_format
    if line is not None:
        if line_rule == 'exact':
            fmt.line_spacing = Pt(line / 20)
        else:
            fmt.line_spacing = line / 240
    if before is not None:
        fmt.space_before = Pt(before)
    if after is not None:
        fmt.space_after = Pt(after)
    if right_indent is not None:
        fmt.right_indent = Pt(right_indent)
    if first_line_indent is not None:
        fmt.first_line_indent = Pt(first_line_indent)
    if page_break_before:
        fmt.page_break_before = True
    return paragraph


def clear_paragraph(paragraph):
    """删除段落里已有的 run，用于复用表格单元格自带的空段落。"""
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    return paragraph


def collapse_cell(cell):
    """合并单元格后只保留一个空段：否则每个被合并格都会留下一个段落撑高行。"""
    for paragraph in list(cell.paragraphs)[1:]:
        paragraph._p.getparent().remove(paragraph._p)
    return cell


# ---------------------------------------------------------------- 表格设置 ----

def build_table(container, widths, *, indent=0, border_size=4, inside_size=None,
                border_color='auto'):
    """建立固定列宽的两列/多列表格，边框与缩进按归档样例设置。

    ``widths`` 为缇为单位的列宽列表；``indent`` 为表格整体左缩进（负值表示向左
    溢出页边距，学院样例的时间安排表就是这样把表格撑得比正文宽）。
    """
    table = container.add_table(rows=0, cols=len(widths))
    table.autofit = False
    _set_table_width(table, sum(widths))
    if indent:
        _set_table_indent(table, indent)
    _set_table_borders(table, border_size, inside_size if inside_size is not None else border_size,
                       border_color)
    for index, width in enumerate(widths):
        table.columns[index].width = Twips(width)
    return table


def add_row(table, *, height=None, height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST, cant_split=False):
    """追加一行，并统一设置行高与单元格垂直居中。"""
    row = table.add_row()
    if height is not None:
        row.height = Twips(height)
        row.height_rule = height_rule
    if cant_split:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split_element = OxmlElement('w:cantSplit')
        tr_pr.insert(0, cant_split_element)
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return row


def set_cell_runs(cell, segments, *, align=None, size=10.5, font=SONG):
    """把若干 ``(文本, 选项)`` 写进单元格首个段落。

    选项字典直接透传给 :func:`styled_run`，用于给同一个人名单独着色。合并单元格
    会留下多个空段，这里统一收成一个，避免把行高撑开。
    """
    collapse_cell(cell)
    paragraph = cell.paragraphs[0]
    clear_paragraph(paragraph)
    if align:
        paragraph.alignment = ALIGNMENTS[align]
    for segment in segments:
        text, options = segment if isinstance(segment, tuple) else (segment, {})
        options = dict(options)
        styled_run(paragraph, text, font=options.pop('font', font),
                   size=options.pop('size', size), **options)
    return paragraph


def set_cell_width(cell, width):
    cell.width = Twips(width)


def merge_row(row, start, end):
    """合并同一行第 ``start`` 到第 ``end`` 列（含），与归档样例的 gridSpan 一致。"""
    return row.cells[start].merge(row.cells[end])


def unique_cells(row):
    """取一行的逻辑单元格。

    python-docx 的 ``row.cells`` 按网格展开合并单元格（一个 tc 会重复出现），
    合并之后按下标写文本会把同一个格反复覆盖，所以写内容一律用它。
    """
    cells, seen = [], set()
    for cell in row.cells:
        if cell._tc in seen:
            continue
        seen.add(cell._tc)
        cells.append(cell)
    return cells


def _set_table_width(table, width):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn('w:tblW'))
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.insert_element_before(tbl_w, *_AFTER_TBL_W)
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_w.set(qn('w:w'), str(int(width)))


def _set_table_indent(table, indent):
    tbl_pr = table._tbl.tblPr
    element = OxmlElement('w:tblInd')
    element.set(qn('w:w'), str(int(indent)))
    element.set(qn('w:type'), 'dxa')
    tbl_pr.insert_element_before(element, *_AFTER_TBL_IND)


def _set_table_borders(table, size, inside_size, color, outer=True):
    tbl_pr = table._tbl.tblPr
    existing = tbl_pr.find(qn('w:tblBorders'))
    if existing is not None:
        tbl_pr.remove(existing)
    borders = OxmlElement('w:tblBorders')
    for edge, edge_size in (('top', size), ('left', size), ('bottom', size), ('right', size),
                            ('insideH', inside_size), ('insideV', inside_size)):
        if not outer and edge in ('top', 'left', 'bottom', 'right'):
            # 不写外框元素即视为继承样式；Normal Table 没有边框，等于不画外框。
            continue
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), str(int(edge_size)))
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)
        borders.append(element)
    tbl_pr.insert_element_before(borders, *_AFTER_TBL_BORDERS)


# ------------------------------------------------------------ 时间段与文案 ----

_SLOT_PATTERN = re.compile(
    r'^\s*(\d{4})-(\d{1,2})-(\d{1,2})\s+(\d{1,2}):(\d{2})\s*[-~－]\s*(\d{1,2}):(\d{2})\s*$')


class Slot:
    """解析后的答辩时间段；无法解析时保留原文，导出不因脏数据失败。"""

    def __init__(self, raw, day=None, start=None, end=None):
        self.raw = (raw or '').strip()
        self.day = day
        self.start = start
        self.end = end

    @property
    def parsed(self):
        return self.day is not None

    @property
    def period(self):
        if not self.parsed:
            return ''
        hour = int(self.start.split(':')[0])
        if hour < 12:
            return '上午'
        return '下午' if hour < 18 else '晚上'

    @property
    def weekday(self):
        return f'周{WEEKDAYS[self.day.weekday()]}' if self.parsed else ''

    @property
    def key(self):
        """单元标识：同一天同一时段视为同一个答辩单元。"""
        return (self.day, self.start, self.end) if self.parsed else (self.raw,)

    def date_text(self):
        return f'{self.day.year}年{self.day.month}月{self.day.day}日'

    def clock_text(self):
        return f'{_clock_text(self.start)}-{_clock_text(self.end)}'


def parse_slot(raw):
    """把 ``2026-04-07 09:00-12:00`` 解析成 :class:`Slot`。"""
    text = (raw or '').strip()
    match = _SLOT_PATTERN.match(text)
    if not match:
        return Slot(text)
    year, month, day, start_hour, start_minute, end_hour, end_minute = match.groups()
    try:
        day_value = date(int(year), int(month), int(day))
    except ValueError:
        return Slot(text)
    return Slot(text, day_value,
                f'{int(start_hour):02d}:{start_minute}', f'{int(end_hour):02d}:{end_minute}')


def arrangement_slot_text(raw):
    """时间安排表格式：``2026年4月1日  周三   上午9:20-12:00``。"""
    slot = parse_slot(raw)
    if not slot.parsed:
        return slot.raw or '待定'
    return f'{slot.date_text()}  {slot.weekday}   {slot.period}{slot.clock_text()}'


def conclusion_slot_text(raw):
    """评审小组结论格式：``2026-04-01（周三）上午9:20-12:00``。"""
    slot = parse_slot(raw)
    if not slot.parsed:
        return slot.raw or '待定'
    return f'{slot.day.isoformat()}（{slot.weekday}）{slot.period}{slot.clock_text()}'


def summary_slot_text(raw):
    """中期考核汇总表格式：``2026年07月13日（周一）上午8:30``（该表用补零日期）。"""
    slot = parse_slot(raw)
    if not slot.parsed:
        return slot.raw or '待定'
    return (f'{slot.day.year}年{slot.day.month:02d}月{slot.day.day:02d}日'
            f'（{slot.weekday}）{slot.period}{_clock_text(slot.start)}')


def slots_of(groups):
    """按组顺序取出去重后的时间段，用于统计单元数。"""
    seen, result = set(), []
    for group in groups:
        slot = parse_slot(getattr(group, 'time', ''))
        if slot.key in seen:
            continue
        seen.add(slot.key)
        result.append(slot)
    return result


def date_span_text(slots):
    """把答辩日期合并成 ``2026年7月13日-16日`` 这样的区间文案。"""
    days = sorted({slot.day for slot in slots if slot.parsed})
    if not days:
        return '待定'
    if len(days) == 1:
        return _date_text(days[0])
    same_year = len({day.year for day in days}) == 1
    same_month = same_year and len({day.month for day in days}) == 1
    if same_month:
        month_days = [day.day for day in days]
        contiguous = month_days == list(range(month_days[0], month_days[-1] + 1))
        head = f'{days[0].year}年{days[0].month}月{days[0].day}日'
        if contiguous:
            return f'{head}-{month_days[-1]}日'
        return '、'.join(_date_text(day) for day in days)
    if same_year:
        head = f'{days[0].year}年{days[0].month}月{days[0].day}日'
        tail = '、'.join(f'{day.month}月{day.day}日' for day in days[1:])
        return f'{head}、{tail}'
    return '、'.join(_date_text(day) for day in days)


def pad_wide(text, width):
    """按“汉字占两格”把姓名补足到固定宽度，让表单里各行的下划线对齐。"""
    text = text or ''
    current = wide_len(text)
    return text + ' ' * max(0, width - current)


def _date_text(day):
    return f'{day.year}年{day.month}月{day.day}日'


def _hour12(hour):
    """13:00 → 1:00；样例写 ``下午2:10-6:00``，小时不补零也不用 24 小时制。"""
    value = int(hour)
    return str(value - 12 if value > 12 else value)


def _clock_text(value):
    hour, minute = value.split(':', 1)
    return f'{_hour12(hour)}:{minute}'


def wide_len(text):
    """显示宽度：汉字和全角符号占两格，西文占一格。"""
    return sum(2 if ord(char) > 127 else 1 for char in text or '')
